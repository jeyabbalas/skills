#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.10"
# dependencies = [
#     "pytest>=8",
#     "openpyxl>=3.1",
#     "pypdf>=4",
#     "cryptography>=3.1",
#     "python-docx>=1.1",
#     "xlrd>=2.0.1",
#     "striprtf>=0.0.26",
#     "xlwt>=1.3",
# ]
# ///
"""Tests for skills/schemify/scripts/extract.py: `uv run tests/test_schemify_extract.py`.

Every input document is generated here (xlwt is a test-only dependency, used to
write the .xls fixture). Fixtures and outputs go under a temporary directory,
or under $EXTRACT_FIXTURES; the subdirectories the tests create
there (src, out, installed, big) are wiped at the start of each session.
"""

import csv
import datetime
import hashlib
import importlib.util
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

sys.dont_write_bytecode = True

HERE = Path(__file__).resolve().parent
SCRIPT = HERE.parent / "skills" / "schemify" / "scripts" / "extract.py"
import tempfile  # noqa: E402
FIXTURES = Path(os.environ.get("EXTRACT_FIXTURES") or tempfile.mkdtemp(prefix="schemify-extract-")).resolve()
SRC = FIXTURES / "src"


# ------------------------------------------------------------------- helpers

def cli(*args):
    """Run extract.py; assert stdout is exactly one JSON object; return (code, obj, stderr)."""
    proc = subprocess.run([sys.executable, str(SCRIPT), *map(str, args)], capture_output=True,
                          text=True, encoding="utf-8")
    assert proc.stdout.lstrip().startswith("{"), proc.stdout + proc.stderr
    obj = json.loads(proc.stdout)  # raises on anything after the object
    assert isinstance(obj, dict)
    assert "Traceback" not in proc.stderr, proc.stderr
    return proc.returncode, obj, proc.stderr


def ok(*args):
    code, obj, err = cli(*args)
    assert code == 0, (obj, err)
    assert obj["ok"] is True
    return obj


def failed(code_expected, *args):
    code, obj, _ = cli(*args)
    assert code == code_expected, obj
    assert set(obj) == {"ok", "error", "hint"}, obj
    assert obj["ok"] is False and obj["error"] and obj["hint"]
    return obj


def fresh(name):
    d = FIXTURES / "out" / name
    shutil.rmtree(d, ignore_errors=True)
    return d


def tree(d):
    return sorted(p.relative_to(d).as_posix() for p in d.rglob("*") if p.is_file())


def lines(path):
    return path.read_text(encoding="utf-8").split("\n")


def csv_rows(path):
    with open(path, encoding="utf-8", newline="") as f:
        return list(csv.reader(f))


def row_by_number(path, n):
    return next(r for r in csv_rows(path)[1:] if r[0] == str(n))


def load_module():
    spec = importlib.util.spec_from_file_location("extract_under_test", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


MANIFEST_KEYS = ["ok", "source", "id", "kind", "sha256", "bytes", "extracted_at", "locators",
                 "files", "files_total", "warnings", "stats"]


def check_manifest(obj, src, kind, out):
    assert list(obj)[:len(MANIFEST_KEYS)] == MANIFEST_KEYS
    assert obj["kind"] == kind
    assert obj["source"] == str(src)
    assert obj["bytes"] == src.stat().st_size
    assert obj["sha256"] == hashlib.sha256(src.read_bytes()).hexdigest()
    assert obj["extracted_at"].endswith("Z")
    datetime.datetime.fromisoformat(obj["extracted_at"].replace("Z", "+00:00"))
    assert isinstance(obj["locators"], str) and obj["locators"]
    disk = json.loads((out / "manifest.json").read_text(encoding="utf-8"))
    assert disk["files"][:50] == obj["files"] and obj["files_total"] == disk["files_total"] == len(disk["files"])
    for entry in disk["files"]:
        assert set(entry) == {"path", "what"} and entry["what"]
        assert (out / entry["path"]).is_file(), entry
    assert tree(out) == sorted([e["path"] for e in disk["files"]] + ["manifest.json"])
    return disk


# ------------------------------------------------------------------ fixtures

def pdf_escape(s):
    return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(path, pages):
    """A minimal hand-written PDF: one Helvetica text block per page ('' = no text layer)."""
    objects = {1: b"<< /Type /Catalog /Pages 2 0 R >>",
               3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica "
                  b"/Encoding /WinAnsiEncoding >>"}
    kids, num = [], 4
    for text in pages:
        ops = ["BT", "/F1 12 Tf", "14 TL", "72 720 Td"]
        ops += [f"({pdf_escape(line)}) Tj T*" for line in text.split("\n")] if text else []
        stream = "\n".join(ops + ["ET"]).encode("cp1252")
        objects[num + 1] = b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream)
        objects[num] = (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                        b"/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % (num + 1))
        kids.append(num)
        num += 2
    objects[2] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (
        b" ".join(b"%d 0 R" % k for k in kids), len(kids))
    out, offsets = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n"), {}
    for i in sorted(objects):
        offsets[i] = len(out)
        out += b"%d 0 obj\n%s\nendobj\n" % (i, objects[i])
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % num
    for i in range(1, num):
        out += b"%010d 00000 n \n" % offsets[i]
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (num, xref)
    path.write_bytes(bytes(out))


def pdf_rewrite(src, dest, outline=(), encrypt=None, labels=False):
    from pypdf import PdfReader, PdfWriter
    writer = PdfWriter(clone_from=PdfReader(str(src)))
    for title, page_index in outline:
        writer.add_outline_item(title, page_index)
    if labels:  # printed page numbers: i, 1, 2
        writer.set_page_label(0, 0, style="/r")
        writer.set_page_label(1, 2, style="/D")
    if encrypt is not None:
        writer.encrypt(user_password=encrypt, owner_password="owner-secret",
                       algorithm="AES-128")
    with open(dest, "wb") as f:
        writer.write(f)


def make_xlsx(path):
    import openpyxl
    from openpyxl.chart import BarChart, Reference
    from openpyxl.comments import Comment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Variables"
    ws["A1"] = "Study X data dictionary"
    ws.merge_cells("A1:C1")
    rows = {3: ["Variable", "Label", "Type", "Updated"],
            4: ["AGE", "Age at enrollment", 12.0, datetime.date(2024, 3, 1)],
            5: ["SEX", "Sex", 3.5, datetime.datetime(2024, 3, 1, 13, 45)],
            6: ["BMI", "Body mass\nindex", 7, None]}
    for r, values in rows.items():
        for c, v in enumerate(values, 1):
            if v is not None:
                ws.cell(row=r, column=c, value=v)
    ws.row_dimensions[5].hidden = True
    ws.column_dimensions["D"].hidden = True
    ws["B4"].comment = Comment("Recoded in 2019", "JB")
    codes = wb.create_sheet("Codes & Values")
    codes.append(["Variable", "Code", "Label"])
    codes["A41"], codes["B41"], codes["C41"] = "SMK", 1, "Current smoker"
    codes["A42"], codes["B42"], codes["C42"] = "SMK", 2, "Former smoker"
    codes["D41"] = "smoking block"
    codes.merge_cells("D41:D42")
    old = wb.create_sheet("Old notes")
    old["A1"] = "deprecated"
    old.sheet_state = "hidden"
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=3, min_row=4, max_row=6))
    wb.create_chartsheet("Chart").add_chart(chart)
    wb.save(path)


def make_xls(path):
    import xlwt
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Variables")
    ws.write_merge(0, 0, 0, 2, "Study X dictionary")
    for c, v in enumerate(["Variable", "Label", "Type", "Updated"]):
        ws.write(2, c, v)
    ws.write(3, 0, "AGE")
    ws.write(3, 1, "Age at enrollment")
    ws.write(3, 2, 12.0)
    ws.write(3, 3, datetime.date(2024, 3, 1), xlwt.easyxf(num_format_str="YYYY-MM-DD"))
    ws.write(4, 0, "SEX")
    ws.write(4, 1, "Sex")
    ws.write(4, 2, 3.5)
    ws.write(4, 3, "hidden column value")
    ws.row(4).hidden = True
    ws.col(3).hidden = True
    codes = wb.add_sheet("Codes")
    codes.write(40, 0, "SMK")
    codes.write(40, 1, 1)
    codes.write(40, 2, "Current smoker")
    old = wb.add_sheet("Old")
    old.write(0, 0, "deprecated")
    old.visibility = 1
    wb.save(str(path))


def make_docx(path):
    import docx
    d = docx.Document()
    d.add_heading("Study X Questionnaire", level=1)                 # ¶1
    d.add_paragraph("Answer every question.")                       # ¶2
    d.add_paragraph("")                                             # ¶3 (empty)
    d.add_heading("Section A", level=2)                             # ¶4
    d.add_paragraph("Do you smoke?", style="List Number")           # ¶5
    d.add_paragraph("How old are you?", style="List Number")        # ¶6
    t = d.add_table(rows=3, cols=3)
    for c, v in enumerate(["Code", "Label", "Note"]):
        t.cell(0, c).text = v
    t.cell(1, 0).text, t.cell(1, 1).text = "1", "Yes"
    t.cell(2, 0).text, t.cell(2, 1).text = "2", "No"
    t.cell(1, 2).merge(t.cell(2, 2)).text = "applies to both"
    d.add_paragraph("After the table.")                             # ¶7
    p = d.add_paragraph("Line one")                                 # ¶8
    p.add_run().add_break()
    p.add_run("line two")
    d.sections[0].header.paragraphs[0].text = "Codebook v2.3"
    d.save(path)


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_RELS = "http://schemas.openxmlformats.org/package/2006/relationships"
OFFICE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
WML = "application/vnd.openxmlformats-officedocument.wordprocessingml"


def make_rich_docx(path):
    """Footnote, content control, legacy checkbox, tracked change, hyperlink, numbering."""
    def run(text):
        return f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r>'

    def numbered(text, level):
        return (f'<w:p><w:pPr><w:numPr><w:ilvl w:val="{level}"/><w:numId w:val="1"/></w:numPr>'
                f'</w:pPr>{run(text)}</w:p>')

    body = "".join([
        f'<w:p>{run("Smoking history")}<w:r><w:footnoteReference w:id="1"/></w:r></w:p>',
        f'<w:sdt><w:sdtPr><w:alias w:val="Q1"/></w:sdtPr><w:sdtContent><w:p>'
        f'{run("Inside a content control")}</w:p></w:sdtContent></w:sdt>',
        '<w:p><w:r><w:fldChar w:fldCharType="begin"><w:ffData><w:checkBox>'
        '<w:default w:val="0"/></w:checkBox></w:ffData></w:fldChar></w:r>'
        '<w:r><w:instrText xml:space="preserve"> FORMCHECKBOX </w:instrText></w:r>'
        f'<w:r><w:fldChar w:fldCharType="end"/></w:r>{run(" Yes")}</w:p>',
        f'<w:p>{run("Kept ")}<w:ins w:id="7" w:author="a">{run("inserted")}</w:ins>'
        '<w:del w:id="8" w:author="a"><w:r><w:delText>deleted</w:delText></w:r></w:del></w:p>',
        f'<w:p><w:hyperlink r:id="rId2">{run("value sets")}</w:hyperlink></w:p>',
        numbered("First question", 0), numbered("Sub a", 1), numbered("Second question", 0),
        f'<w:p><w:pPr><w:outlineLvl w:val="2"/></w:pPr>{run("Outline heading")}</w:p>',
        "<w:sectPr/>"])
    parts = {
        "[Content_Types].xml":
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.'
            'relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
            f'<Override PartName="/word/document.xml" ContentType="{WML}.document.main+xml"/>'
            f'<Override PartName="/word/footnotes.xml" ContentType="{WML}.footnotes+xml"/>'
            f'<Override PartName="/word/numbering.xml" ContentType="{WML}.numbering+xml"/>'
            '</Types>',
        "_rels/.rels":
            f'<Relationships xmlns="{PKG_RELS}"><Relationship Id="rId1" '
            f'Type="{OFFICE_REL}/officeDocument" Target="word/document.xml"/></Relationships>',
        "word/_rels/document.xml.rels":
            f'<Relationships xmlns="{PKG_RELS}">'
            f'<Relationship Id="rId1" Type="{OFFICE_REL}/footnotes" Target="footnotes.xml"/>'
            f'<Relationship Id="rId2" Type="{OFFICE_REL}/hyperlink" '
            'Target="https://example.org/valuesets" TargetMode="External"/>'
            f'<Relationship Id="rId3" Type="{OFFICE_REL}/numbering" Target="numbering.xml"/>'
            '</Relationships>',
        "word/document.xml":
            f'<w:document xmlns:w="{W_NS}" xmlns:r="{R_NS}"><w:body>{body}</w:body></w:document>',
        "word/numbering.xml":
            f'<w:numbering xmlns:w="{W_NS}"><w:abstractNum w:abstractNumId="0">'
            '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
            '<w:lvlText w:val="Q%1."/></w:lvl><w:lvl w:ilvl="1"><w:start w:val="1"/>'
            '<w:numFmt w:val="lowerLetter"/><w:lvlText w:val="%2)"/></w:lvl></w:abstractNum>'
            '<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num></w:numbering>',
        "word/footnotes.xml":
            f'<w:footnotes xmlns:w="{W_NS}">'
            '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p>'
            '</w:footnote><w:footnote w:id="1"><w:p><w:r><w:footnoteRef/></w:r>'
            f'{run(" Includes e-cigarettes.")}</w:p></w:footnote></w:footnotes>',
    }
    import zipfile
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, xml in parts.items():
            z.writestr(name, '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n' + xml)


HTML = """<!DOCTYPE html>
<html><head><meta charset="utf-8"><title>Study X codebook</title>
<style>p { color: red }</style><script>var secret = "DO-NOT-EXTRACT";</script></head>
<body>
<h1>Codebook</h1>
<p>See the <a href="https://example.org/docs">documentation</a> for details.</p>
<h2>Variables</h2>
<table>
<tr><th>Name</th><th>Label</th></tr>
<tr><td>AGE</td><td>&Acirc;ge<br>in years</td></tr>
</table>
<ul><li>First item</li><li>Second item</li></ul>
</body></html>
"""

RTF = (r"{\rtf1\ansi\ansicpg1252\deff0{\fonttbl{\f0 Helvetica;}}\f0\fs24 Codebook\par "
       r"Variable AGE: \'c2ge du participant\par Line three\par}")

CSV_CP1252 = ("Variable;Label;Type\r\n"
              "AGE;Âge à l'inclusion;num\r\n"
              "SEX;Sexe;\"cat; 1=M; 2=F\"\r\n"
              "\r\n"
              "BMI;\"Indice de masse\r\ncorporelle\";num\r\n"
              "WT;Poids;num\r\n")

PDF_PAGES = ["Study X protocol\nVariable AGE: age at enrollment in years",
             "Codes for smoking\nSMK 1 = current smoker; 2 = former smoker",
             ""]


@pytest.fixture(scope="session", autouse=True)
def src():
    for name in ("src", "out", "installed", "big"):
        shutil.rmtree(FIXTURES / name, ignore_errors=True)
    SRC.mkdir(parents=True)
    make_xlsx(SRC / "dictionary.xlsx")
    make_xls(SRC / "legacy.xls")
    make_docx(SRC / "questionnaire.docx")
    (SRC / "codes.csv").write_bytes(CSV_CP1252.encode("cp1252"))
    (SRC / "vars.txt").write_text("Variable\tLabel\tType\nAGE\tAge\tnum\nSEX\tSex\tcat\n"
                                  "BMI\tBody mass index\tnum\nWT\tWeight\tnum\n",
                                  encoding="utf-8")
    (SRC / "notes.md").write_bytes(b"# Notes\r\n\r\nLine three\r\nLine four")
    (SRC / "page.html").write_text(HTML, encoding="utf-8")
    (SRC / "codebook.rtf").write_text(RTF, encoding="ascii")
    (SRC / "schema.json").write_text(json.dumps(
        {"name": "study", "variables": [{"name": "AGE", "label": "Âge"}]}), encoding="utf-8")
    (SRC / "broken.json").write_text('{"name": "study",\n"oops"\n}', encoding="utf-8")
    (SRC / "ddi.xml").write_text(
        '<?xml version="1.0" encoding="UTF-8"?><codeBook xmlns="ddi:codebook:2_5"><dataDscr>'
        '<var name="AGE"><labl>Age</labl></var></dataDscr></codeBook>', encoding="utf-8")
    make_pdf(SRC / "plain.pdf", PDF_PAGES)
    pdf_rewrite(SRC / "plain.pdf", SRC / "protocol.pdf",
                outline=[("Introduction", 0), ("Codes", 1)])
    make_pdf(SRC / "long.pdf", [f"Page {i} describes variable V{i:03d} in detail."
                                for i in range(1, 61)])
    pdf_rewrite(SRC / "plain.pdf", SRC / "locked.pdf", encrypt="secret")
    pdf_rewrite(SRC / "plain.pdf", SRC / "owner-only.pdf", encrypt="")
    (SRC / "data.sav").write_bytes(b"$FL2@(#) IBM SPSS STATISTICS" + bytes(64))
    (SRC / "renamed.txt").write_bytes(b"SQLite format 3\x00" + bytes(100))
    (SRC / "mystery.xyz").write_text("what is this", encoding="utf-8")
    (SRC / "old.doc").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(504))
    (SRC / "webpage.xls").write_text("<html><body><table><tr><td>AGE</td><td>Age</td></tr>"
                                     "</table></body></html>", encoding="utf-8")
    (SRC / "empty.txt").write_bytes(b"")
    make_rich_docx(SRC / "rich.docx")
    pdf_rewrite(SRC / "plain.pdf", SRC / "labelled.pdf", labels=True)
    unicode_text = "".join(f"V{i}\tLabel {i} \u00e9\tnum\r\n" for i in range(1, 7))
    (SRC / "unicode.txt").write_bytes(b"\xff\xfe" + ("Variable\tLabel\tType\r\n"
                                                     + unicode_text).encode("utf-16-le"))
    (SRC / "sloppy.html").write_text(
        "<title>Unclosed title<table><tr><td>A<td>B<tr><td>C<td>D</table>"
        "<ul><li>one<li>two<br>(go to Q5)</ul><p>para one<p>para two", encoding="utf-8")
    (SRC / "ragged.csv").write_text("a,b,c\n1,2,3\n4,5\n6,7,8\n", encoding="utf-8")
    (SRC / "locked.xlsx").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(1000)
                                      + "EncryptionInfo".encode("utf-16-le") + bytes(1000))
    return SRC


# --------------------------------------------------------------------- tests

def test_help_documents_usage_and_exit_codes():
    proc = subprocess.run([sys.executable, str(SCRIPT), "--help"], capture_output=True,
                          text=True, encoding="utf-8")
    assert proc.returncode == 0
    for needle in ("--out OUTDIR", "--max-pages", "--id", "--documentation", "Exit codes", "python3 -m pip install",
                   "uv run extract.py", "sheet-NN-<slug>.csv", "pages/pNNNN.txt", "¶N"):
        assert needle in proc.stdout, needle


@pytest.mark.parametrize("args", [
    [],
    ["SRC/dictionary.xlsx"],
    ["SRC/missing.xlsx", "--out", "OUT"],
    ["SRC/dictionary.xlsx", "--out", "OUT", "--max-pages", "0"],
    ["SRC/dictionary.xlsx", "--out", "OUT", "--max-pages", "many"],
    ["SRC", "--out", "OUT"],
    ["SRC/empty.txt", "--out", "OUT"],
])
def test_usage_and_input_errors_exit_1(args):
    out = fresh("usage")
    args = [a.replace("SRC", str(SRC)).replace("OUT", str(out)) for a in args]
    failed(1, *args)
    assert not out.exists()


def test_xlsx(src):
    out = fresh("xlsx")
    source = src / "dictionary.xlsx"
    obj = ok(source, "--out", out, "--id", "S01")
    disk = check_manifest(obj, source, "xlsx", out)
    assert obj["id"] == "S01"
    assert [f["path"] for f in disk["files"]] == [
        "sheet-01-variables.csv", "sheet-02-codes-values.csv", "sheet-03-old-notes.csv",
        "sheets.json", "comments.csv"]
    assert disk["files"][0]["what"] == "sheet 'Variables' (5 rows x 4 cols)"
    variables = out / "sheet-01-variables.csv"
    assert csv_rows(variables)[0] == ["_row", "A", "B", "C", "D"]
    assert [r[0] for r in csv_rows(variables)[1:]] == ["1", "3", "4", "5", "6"]  # row 2 empty
    assert row_by_number(variables, 1) == ["1"] + ["Study X data dictionary"] * 3 + [""]
    assert row_by_number(variables, 4) == ["4", "AGE", "Age at enrollment", "12", "2024-03-01"]
    assert row_by_number(variables, 5) == ["5", "SEX", "Sex", "3.5", "2024-03-01T13:45:00"]
    assert row_by_number(variables, 6) == ["6", "BMI", "Body mass\\nindex", "7", ""]
    assert len(lines(variables)) == 7  # header + 5 rows + final newline: one line per row
    codes = out / "sheet-02-codes-values.csv"
    assert row_by_number(codes, 41) == ["41", "SMK", "1", "Current smoker", "smoking block"]
    assert row_by_number(codes, 42) == ["42", "SMK", "2", "Former smoker", "smoking block"]
    sheets = obj["stats"]["sheets"]
    assert json.loads((out / "sheets.json").read_text(encoding="utf-8")) == sheets
    first = sheets[0]
    assert (first["name"], first["state"], first["dims"]) == ("Variables", "visible", "A1:D6")
    assert (first["first_nonempty_row"], first["header_guess"]) == (1, 3)
    assert (first["rows"], first["cols"]) == (5, 4)
    assert first["merged"] == ["A1:C1"]
    assert (first["hidden_rows"], first["hidden_cols"]) == (["5"], ["D"])
    assert sheets[1]["merged"] == ["D41:D42"] and sheets[1]["header_guess"] == 1
    assert sheets[2]["state"] == "hidden"
    warnings = " ".join(obj["warnings"])
    assert "sheet 'Old notes' is hidden" in warnings
    assert "row 5 is hidden" in warnings and "column D is hidden" in warnings
    assert csv_rows(out / "comments.csv")[1] == ["Variables", "B4", "4", "JB", "Recoded in 2019"]
    assert "sheet 4 'Chart' is a chart sheet" in warnings and len(sheets) == 3
    assert "_row is the Excel row number" in obj["locators"]


def test_xls(src):
    out = fresh("xls")
    source = src / "legacy.xls"
    obj = ok(source, "--out", out)
    check_manifest(obj, source, "xls", out)
    assert obj["id"] is None
    variables = out / "sheet-01-variables.csv"
    assert row_by_number(variables, 1)[1:4] == ["Study X dictionary"] * 3
    assert row_by_number(variables, 4) == ["4", "AGE", "Age at enrollment", "12", "2024-03-01"]
    assert row_by_number(out / "sheet-02-codes.csv", 41) == ["41", "SMK", "1", "Current smoker"]
    sheets = obj["stats"]["sheets"]
    assert sheets[0]["merged"] == ["A1:C1"] and sheets[0]["header_guess"] == 3
    assert (sheets[0]["hidden_rows"], sheets[0]["hidden_cols"]) == (["5"], ["D"])
    assert sheets[2]["state"] == "hidden"
    assert any("'Old' is hidden" in w for w in obj["warnings"])


def test_csv_cp1252_semicolon(src):
    out = fresh("csv")
    source = src / "codes.csv"
    obj = ok(source, "--out", out, "--documentation")
    check_manifest(obj, source, "csv", out)
    stats = obj["stats"]
    assert (stats["encoding"], stats["delimiter"]) == ("cp1252", ";")
    assert (stats["rows"], stats["columns"], stats["header_line"]) == (4, 3, 1)
    table = out / "table.csv"
    assert csv_rows(table)[0] == ["_row", "Variable", "Label", "Type"]
    assert row_by_number(table, 2) == ["2", "AGE", "Âge à l'inclusion", "num"]
    assert row_by_number(table, 3) == ["3", "SEX", "Sexe", "cat; 1=M; 2=F"]
    assert row_by_number(table, 5) == ["5", "BMI", "Indice de masse\\ncorporelle", "num"]
    assert row_by_number(table, 7) == ["7", "WT", "Poids", "num"]
    assert table.read_bytes().decode("utf-8")  # re-encoded as UTF-8


def test_delimited_txt_and_markdown(src):
    out = fresh("txt")
    obj = ok(src / "vars.txt", "--out", out, "--documentation")
    check_manifest(obj, src / "vars.txt", "txt", out)
    assert [f["path"] for f in obj["files"]] == ["text.md", "table.csv"]
    assert obj["stats"]["delimiter"] == "\t" and obj["stats"]["lines"] == 5
    assert row_by_number(out / "table.csv", 4) == ["4", "BMI", "Body mass index", "num"]
    assert lines(out / "text.md")[3] == "BMI\tBody mass index\tnum"  # L4
    out = fresh("md")
    obj = ok(src / "notes.md", "--out", out)
    assert [f["path"] for f in obj["files"]] == ["text.md"]
    assert (out / "text.md").read_bytes() == b"# Notes\n\nLine three\nLine four\n"
    assert obj["stats"]["lines"] == 4 and "L<N>" in obj["locators"]


def test_pdf(src):
    out = fresh("pdf")
    source = src / "protocol.pdf"
    obj = ok(source, "--out", out)
    check_manifest(obj, source, "pdf", out)
    assert [f["path"] for f in obj["files"]] == [
        "pages/p0001.txt", "pages/p0002.txt", "pages/p0003.txt", "text.md"]
    stats = obj["stats"]
    assert (stats["pages"], stats["pages_extracted"], stats["empty_pages"]) == (3, 3, [3])
    assert stats["outline"] == [{"title": "Introduction", "page": 1, "level": 1},
                                {"title": "Codes", "page": 2, "level": 1}]
    assert "SMK 1 = current smoker" in (out / "pages/p0002.txt").read_text(encoding="utf-8")
    text = lines(out / "text.md")
    markers = [i for i, line in enumerate(text) if line.startswith("<!-- page ")]
    assert [text[i] for i in markers] == ["<!-- page 1 -->", "<!-- page 2 -->", "<!-- page 3 -->"]
    page2 = "\n".join(text[markers[1] + 1:markers[2]])
    assert "Codes for smoking" in page2 and "Study X protocol" not in page2
    assert any("view these pages directly; the text layer is empty" in w and "page 3 is" in w
               for w in obj["warnings"])
    assert "p.N" in obj["locators"]


def test_pdf_max_pages_then_full_leaves_no_leftovers(src):
    out = fresh("pdf-max")
    ok(src / "protocol.pdf", "--out", out)
    obj = ok(src / "protocol.pdf", "--out", out, "--max-pages", "1")
    assert tree(out) == ["manifest.json", "pages/p0001.txt", "text.md"]
    assert (obj["stats"]["pages"], obj["stats"]["pages_extracted"]) == (3, 1)
    assert any("--max-pages 1" in w for w in obj["warnings"])
    ok(src / "protocol.pdf", "--out", out)
    assert len(tree(out)) == 5


def test_stdout_is_bounded(src):
    out = fresh("long-pdf")
    obj = ok(src / "long.pdf", "--out", out)
    assert len(obj["files"]) == 50 and obj["files_total"] == 61
    disk = json.loads((out / "manifest.json").read_text(encoding="utf-8"))
    assert len(disk["files"]) == 61 and disk["files_total"] == 61
    assert disk["files"][-1]["path"] == "text.md"
    assert "V060" in (out / "pages/p0060.txt").read_text(encoding="utf-8")


def test_encrypted_pdfs(src):
    failed(2, src / "locked.pdf", "--out", fresh("locked"))
    out = fresh("owner-only")
    obj = ok(src / "owner-only.pdf", "--out", out)
    assert obj["stats"]["encrypted"] is True
    assert "Study X protocol" in (out / "pages/p0001.txt").read_text(encoding="utf-8")


def test_docx(src):
    out = fresh("docx")
    source = src / "questionnaire.docx"
    obj = ok(source, "--out", out)
    check_manifest(obj, source, "docx", out)
    assert [f["path"] for f in obj["files"]] == ["table-01.csv", "text.md", "notes.md"]
    assert lines(out / "text.md") == [
        "¶1 # Study X Questionnaire",
        "¶2 Answer every question.",
        "¶4 ## Section A",
        "¶5 1. Do you smoke?",
        "¶6 2. How old are you?",
        "[table-01.csv: 3 rows x 3 cols]",
        "¶7 After the table.",
        "¶8 Line one\\nline two",
        ""]
    assert csv_rows(out / "table-01.csv") == [
        ["_row", "A", "B", "C"], ["1", "Code", "Label", "Note"],
        ["2", "1", "Yes", "applies to both"], ["3", "2", "No", "applies to both"]]
    assert obj["files"][0]["what"] == "table 1 (3 rows x 3 cols), after ¶6"
    stats = obj["stats"]
    assert (stats["paragraphs"], stats["tables"], stats["headings"]) == (8, 1, 2)
    assert stats["list_items"] == 2 and stats["header_footer_lines"] == 1
    assert "Codebook v2.3" in (out / "notes.md").read_text(encoding="utf-8")
    assert any("automatic list numbering" in w for w in obj["warnings"])


def test_html(src):
    out = fresh("html")
    source = src / "page.html"
    obj = ok(source, "--out", out)
    check_manifest(obj, source, "html", out)
    text = lines(out / "text.md")
    assert text[0] == "# Codebook"
    assert "See the documentation (https://example.org/docs) for details." in text
    assert text.index("## Variables") < text.index("Name | Label") < text.index(
        "AGE | Âge\\nin years")
    assert "- First item" in text and "- Second item" in text
    assert "DO-NOT-EXTRACT" not in "\n".join(text) and "color: red" not in "\n".join(text)
    stats = obj["stats"]
    assert (stats["links"], stats["tables"], stats["headings"]) == (1, 1, 2)
    assert stats["title"] == "Study X codebook" and stats["chars"] > 0


def test_rtf(src):
    out = fresh("rtf")
    source = src / "codebook.rtf"
    obj = ok(source, "--out", out)
    check_manifest(obj, source, "rtf", out)
    text = (out / "text.md").read_text(encoding="utf-8")
    assert "Âge du participant" in text and text.splitlines()[0].strip() == "Codebook"
    expected = "textutil" if shutil.which("textutil") else "striprtf"
    assert obj["stats"]["converter"] == expected and obj["stats"]["lines"] >= 3


def test_rtf_without_textutil_uses_striprtf(src, monkeypatch, capsys):
    module = load_module()
    monkeypatch.setattr(module.shutil, "which", lambda name: None)
    out = fresh("rtf-striprtf")
    assert module.main([str(src / "codebook.rtf"), "--out", str(out)]) == 0
    obj = json.loads(capsys.readouterr().out)
    assert obj["stats"]["converter"] == "striprtf"
    text = (out / "text.md").read_text(encoding="utf-8")
    assert "Âge du participant" in text and "Line three" in text


def test_json_and_xml(src):
    out = fresh("json")
    obj = ok(src / "schema.json", "--out", out)
    check_manifest(obj, src / "schema.json", "json", out)
    assert obj["stats"]["type"] == "object" and obj["stats"]["keys"] == ["name", "variables"]
    pretty = (out / "pretty.json").read_text(encoding="utf-8")
    assert '"label": "Âge"' in pretty and json.loads(pretty)["variables"][0]["name"] == "AGE"
    out = fresh("broken-json")
    obj = ok(src / "broken.json", "--out", out)
    assert [f["path"] for f in obj["files"]] == ["text.md"]
    assert any("not valid JSON" in w for w in obj["warnings"])
    out = fresh("xml")
    obj = ok(src / "ddi.xml", "--out", out)
    check_manifest(obj, src / "ddi.xml", "xml", out)
    stats = obj["stats"]
    assert (stats["root"], stats["namespace"], stats["elements"]) == (
        "codeBook", "ddi:codebook:2_5", 4)
    assert "    <var name=\"AGE\">" in lines(out / "pretty.xml")


def test_mislabeled_xls_is_routed_by_content(src):
    out = fresh("webpage")
    obj = ok(src / "webpage.xls", "--out", out)
    assert obj["kind"] == "html"
    assert any("content is html" in w for w in obj["warnings"])
    assert "AGE | Age" in lines(out / "text.md")


def test_data_files_are_refused(src):
    obj = failed(2, src / "data.sav", "--out", fresh("sav"))
    assert obj["error"] == ("This looks like a data file, not a metadata document. "
                            "The agent must not open data files.")
    assert obj["hint"].startswith("If it embeds variable/value labels you need")
    obj = failed(2, src / "renamed.txt", "--out", fresh("renamed"))
    assert obj["error"].startswith("This looks like a data file") and "SQLite" in obj["error"]
    assert not (FIXTURES / "out" / "sav").exists()


def test_unsupported_formats_exit_2(src):
    for name in ("mystery.xyz", "old.doc"):
        obj = failed(2, src / name, "--out", fresh(name))
        assert "Supported: .xlsx" in obj["hint"] and ".pdf" in obj["hint"]
    assert ".txt extension" in failed(2, src / "mystery.xyz", "--out", fresh("x"))["hint"]
    obj = failed(2, src / "locked.xlsx", "--out", fresh("locked-xlsx"))
    assert "password-protected" in obj["error"]


def test_size_guard(src):
    big = FIXTURES / "big" / "huge.txt"
    big.parent.mkdir(parents=True, exist_ok=True)
    with open(big, "wb") as f:
        f.truncate(201 * 1024 * 1024)  # sparse: no real disk use
    try:
        obj = failed(2, big, "--out", fresh("big"))
        assert "200 MB" in obj["error"]
    finally:
        big.unlink()


def test_refuses_to_write_inside_schemify_skill(src):
    skill = FIXTURES / "installed" / "schemify"
    (skill / "scripts").mkdir(parents=True, exist_ok=True)
    (skill / "SKILL.md").write_text('---\nname: "schemify"\ndescription: "x"\n---\nBody\n',
                                    encoding="utf-8")
    target = skill / "scripts" / "extract-out"
    obj = failed(1, src / "codes.csv", "--out", target, "--documentation")
    assert "skill" in obj["error"] and "--out" in obj["hint"]
    assert not target.exists()
    other = FIXTURES / "installed" / "other"
    other.mkdir(parents=True, exist_ok=True)
    (other / "SKILL.md").write_text("---\nname: other-skill\n---\n", encoding="utf-8")
    ok(src / "codes.csv", "--out", other / "out", "--documentation")  # another skill's tree is fine


def test_rerun_is_idempotent_and_spares_foreign_files(src):
    out = fresh("idem")
    first = ok(src / "dictionary.xlsx", "--out", out)
    snapshot = {p: (out / p).read_bytes() for p in tree(out) if p != "manifest.json"}
    (out / "my-notes.txt").write_text("keep me", encoding="utf-8")
    second = ok(src / "dictionary.xlsx", "--out", out)
    assert {p: (out / p).read_bytes() for p in tree(out)
            if p not in ("manifest.json", "my-notes.txt")} == snapshot
    strip = ("extracted_at",)
    assert {k: v for k, v in first.items() if k not in strip} == {
        k: v for k, v in second.items() if k not in strip}
    assert (out / "my-notes.txt").read_text(encoding="utf-8") == "keep me"
    ok(src / "protocol.pdf", "--out", out)  # a different source into the same directory
    obj = ok(src / "codes.csv", "--out", out, "--documentation")
    assert tree(out) == ["manifest.json", "my-notes.txt", "table.csv"]
    assert any("different source" in w for w in obj["warnings"])


def test_never_overwrites_files_it_did_not_write(src):
    out = fresh("foreign")
    out.mkdir(parents=True)
    (out / "text.md").write_text("the steward's own notes", encoding="utf-8")
    failed(1, src / "notes.md", "--out", out)
    assert (out / "text.md").read_text(encoding="utf-8") == "the steward's own notes"
    out = fresh("foreign-manifest")
    out.mkdir(parents=True)
    (out / "manifest.json").write_text('{"name": "someone else"}', encoding="utf-8")
    failed(1, src / "notes.md", "--out", out)
    assert tree(out) == ["manifest.json"]


def test_source_inside_its_own_output_is_refused(src):
    out = fresh("self")
    ok(src / "notes.md", "--out", out)
    failed(1, out / "text.md", "--out", out)
    assert (out / "text.md").is_file()


def test_large_workbooks_are_streamed(src, monkeypatch, capsys):
    module = load_module()
    monkeypatch.setattr(module, "XLSX_STREAM_XML", 0)
    out = fresh("xlsx-streamed")
    assert module.main([str(src / "dictionary.xlsx"), "--out", str(out)]) == 0
    obj = json.loads(capsys.readouterr().out)
    assert obj["stats"]["mode"] == "streamed"
    assert any("streamed" in w for w in obj["warnings"])
    variables = out / "sheet-01-variables.csv"
    assert row_by_number(variables, 1) == ["1", "Study X data dictionary", "", "", ""]
    assert row_by_number(variables, 4) == ["4", "AGE", "Age at enrollment", "12", "2024-03-01"]
    assert row_by_number(variables, 6) == ["6", "BMI", "Body mass\\nindex", "7", ""]
    codes = out / "sheet-02-codes-values.csv"
    assert row_by_number(codes, 41) == ["41", "SMK", "1", "Current smoker", "smoking block"]
    sheets = obj["stats"]["sheets"]
    assert (sheets[0]["dims"], sheets[0]["header_guess"], sheets[0]["rows"]) == ("A1:D6", 3, 5)
    assert sheets[2]["state"] == "hidden" and sheets[0]["merged"] == []


def test_docx_features_python_docx_cannot_write(src):
    out = fresh("rich-docx")
    obj = ok(src / "rich.docx", "--out", out)
    assert lines(out / "text.md") == [
        "¶1 Smoking history[^1]",
        "¶2 Inside a content control",
        "¶3 ☐ Yes",
        "¶4 Kept inserted",
        "¶5 value sets (https://example.org/valuesets)",
        "¶6 Q1. First question",
        "¶7 a) Sub a",
        "¶8 Q2. Second question",
        "¶9 ### Outline heading",
        ""]
    assert "[^1] Includes e-cigarettes." in lines(out / "notes.md")
    assert obj["stats"]["footnotes"] == 1 and obj["stats"]["list_items"] == 3
    joined = " ".join(obj["warnings"])
    assert "tracked changes" in joined and "3 paragraphs use Word's automatic" in joined


def test_pdf_page_labels(src):
    obj = ok(src / "labelled.pdf", "--out", fresh("labelled"))
    assert obj["stats"]["page_labels"] == ["i", "1", "2"]
    assert any("physical p.1 is labelled 'i'" in w for w in obj["warnings"])


def test_utf16_tab_text_from_excel(src):
    out = fresh("unicode")
    obj = ok(src / "unicode.txt", "--out", out, "--documentation")
    assert obj["stats"]["encoding"] == "utf-16" and obj["stats"]["delimiter"] == "\t"
    assert row_by_number(out / "table.csv", 3) == ["3", "V2", "Label 2 é", "num"]
    assert lines(out / "text.md")[2] == "V2\tLabel 2 é\tnum"


def test_html_with_implied_end_tags(src):
    out = fresh("sloppy")
    obj = ok(src / "sloppy.html", "--out", out)
    got = [ln for ln in lines(out / "text.md") if ln]
    assert got[-7:] == [
        "A | B", "C | D", "- one", "- two", "  (go to Q5)", "para one", "para two"]
    # An unclosed <title>: older parsers end it at the next tag, HTML5 ones read to EOF.
    assert got[:-7] == ["Unclosed title"] or obj["stats"]["title"] == "Unclosed title"


def test_delimited_files_need_documentation_flag(src):
    # A delimited file is as often a dataset as a dictionary: refused (exit 2) until the
    # user has said it describes variables, and nothing is written.
    for name in ("codes.csv", "vars.txt", "unicode.txt"):
        out = fresh("guard")
        obj = failed(2, src / name, "--out", out)
        assert "--documentation" in obj["hint"]
        assert not out.exists()
    out = fresh("guard-md")
    ok(src / "notes.md", "--out", out)  # plain text needs no flag


def test_sheets_of_records_are_not_extracted(src):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Variables"
    ws.append(["Variable", "Label", "Type"])
    for row in (["age", "Age at baseline", "integer"], ["sex", "Sex", "integer"], ["bmi", "Body mass index", "number"],
                ["smk_ever", "Ever smoked", "integer"]):
        ws.append(row)
    data = wb.create_sheet("Sample data")
    data.append(["id", "age", "sex", "bmi", "smk_ever"])
    for k in range(6):
        data.append([f"P{k:03d}", 40 + k, 1 + k % 2, 22.5 + k, 2])
    data["B3"].comment = __import__("openpyxl.comments", fromlist=["Comment"]).Comment("entered twice", "JB")
    loose = wb.create_sheet("Export")                     # records with no dictionary sheet naming its columns
    loose.append(["pid", "visit", "sbp", "dbp"])
    for k in range(8):
        loose.append([k, 1, 120 + k, 80 + k])
    path = src / "dictionary-with-data.xlsx"
    wb.save(path)

    out = fresh("xlsx-records")
    obj = ok(path, "--out", out)
    sheets = {x["name"]: x for x in obj["stats"]["sheets"]}
    assert sheets["Variables"]["file"] == "sheet-01-variables.csv"
    assert "another sheet lists" in sheets["Sample data"]["skipped"] and sheets["Sample data"]["file"] is None
    assert "mostly numbers" in sheets["Export"]["skipped"]
    assert "rows" not in sheets["Sample data"]                    # not even measured
    assert obj["stats"]["skipped_sheets"] == ["Sample data", "Export"]
    assert sorted(p.name for p in out.glob("*.csv")) == ["sheet-01-variables.csv"]   # its comment is dropped too
    assert any("--include-sheet \"Sample data\"" in w for w in obj["warnings"])

    out2 = fresh("xlsx-records-included")
    obj = ok(path, "--out", out2, "--include-sheet", "Sample data")
    assert (out2 / "sheet-02-sample-data.csv").is_file() and obj["stats"]["skipped_sheets"] == ["Export"]


def test_cleaning_code_is_text_with_embedded_data_blanked(src):
    sas = src / "clean.sas"
    sas.write_text("data codes;\n  input code label $;\n  datalines;\n1 Yes\n2 No\n;\nrun;\n"
                   "if smk_ever = 2 then cigs_day = -6;\n", encoding="utf-8")
    out = fresh("code-sas")
    obj = ok(sas, "--out", out)
    assert obj["kind"] == "code" and obj["stats"]["embedded_data_lines"] == 2
    text = lines(out / "text.md")
    assert text[2] == "  datalines;" and text[3] == text[4] == "[embedded data line, not extracted]"
    assert text[7] == "if smk_ever = 2 then cigs_day = -6;"            # line N is still source line N
    assert any(w.startswith("L4-L5:") for w in obj["warnings"])
    do = src / "clean.do"
    do.write_text("input id age\n1 40\nend\nreplace x = -6 if y == 2\n", encoding="utf-8")
    out = fresh("code-do")
    obj = ok(do, "--out", out)
    assert lines(out / "text.md")[1] == "[embedded data line, not extracted]" and obj["stats"]["embedded_data_lines"] == 1
    r = src / "derive.R"
    r.write_text("bmi <- weight / height^2\n", encoding="utf-8")
    assert ok(r, "--out", fresh("code-r"))["stats"]["embedded_data_lines"] == 0


def test_ragged_csv_is_reported(src):
    out = fresh("ragged")
    obj = ok(src / "ragged.csv", "--out", out, "--documentation")
    assert obj["stats"]["ragged_rows"] == 1
    assert any("1 row has a different number of fields" in w and "line 3" in w
               for w in obj["warnings"])
    assert row_by_number(out / "table.csv", 3) == ["3", "4", "5"]


def test_huge_xml_is_copied_not_pretty_printed(src, monkeypatch, capsys):
    module = load_module()
    monkeypatch.setattr(module, "XML_PRETTY_MAX", 0)
    out = fresh("xml-raw")
    assert module.main([str(src / "ddi.xml"), "--out", str(out)]) == 0
    obj = json.loads(capsys.readouterr().out)
    assert [f["path"] for f in obj["files"]] == ["source.xml"]
    assert obj["stats"]["pretty"] is False and obj["stats"]["elements"] == 4
    assert "source.xml line N" in obj["locators"]


def test_missing_dependency_is_a_clean_error(src, monkeypatch, capsys):
    module = load_module()
    monkeypatch.setitem(sys.modules, "openpyxl", None)
    out = fresh("nodeps")
    assert module.main([str(src / "dictionary.xlsx"), "--out", str(out)]) == 1
    obj = json.loads(capsys.readouterr().out)
    assert set(obj) == {"ok", "error", "hint"} and obj["ok"] is False
    assert "python3 -m pip install --user openpyxl pypdf python-docx xlrd striprtf" in obj["hint"]
    assert not out.exists()


def test_value_formatting_helpers():
    module = load_module()
    assert module.cell_text(12.0) == "12" and module.cell_text(0.1) == "0.1"
    assert module.cell_text(1e20) == "1e+20" and module.cell_text(True) == "TRUE"
    assert module.cell_text(datetime.datetime(2024, 3, 1)) == "2024-03-01"
    assert module.cell_text("a\r\nb") == "a\\nb" and module.cell_text("   ") == ""
    assert module.slug("Codes & Values (2019)") == "codes-values-2019"
    assert module.slug("Données") == "donnees" and module.slug("変数") == "sheet"
    assert module.ranges([1, 2, 3, 7, 9, 10]) == ["1-3", "7", "9-10"]
    assert [module.col_letter(n) for n in (1, 26, 27, 702, 703)] == ["A", "Z", "AA", "ZZ", "AAA"]


def test_uv_resolves_the_inline_metadata():
    uv = shutil.which("uv")
    if not uv:
        pytest.skip("uv is not installed")
    proc = subprocess.run([uv, "run", "--offline", "--script", str(SCRIPT), "--help"],
                          capture_output=True, text=True, encoding="utf-8")
    if proc.returncode != 0 and ("offline" in proc.stderr.lower()
                                 or "network" in proc.stderr.lower()):
        pytest.skip("dependencies are not in uv's cache and the test runs offline")
    assert proc.returncode == 0, proc.stderr
    assert "Exit codes" in proc.stdout


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q", "-p", "no:cacheprovider"]))
