#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.10"
# dependencies = [
#     "openpyxl>=3.1",
#     "pypdf>=4",
#     "cryptography>=3.1",
#     "python-docx>=1.1",
#     "xlrd>=2.0.1",
#     "striprtf>=0.0.26",
# ]
# ///
"""Extract one metadata document into plain text/CSV files with stable locators.

Data dictionaries, codebooks, questionnaires and protocols arrive as
spreadsheets, PDFs, Word files and web pages. This script turns ONE of them,
deterministically, into UTF-8 text and CSV files that can be grepped and cited
exactly ("sheet Codes row 41", "p.6", "¶12", "L120"). Do not parse these
formats by hand: run this and work from its output.

Run with uv, which installs the dependencies:
    uv run extract.py SOURCE --out OUTDIR [--id S01] [--max-pages N] [--documentation]
Without uv:
    python3 -m pip install --user openpyxl pypdf python-docx xlrd striprtf
    python3 extract.py SOURCE --out OUTDIR
(add `cryptography` to that pip line for AES-encrypted PDFs).

Arguments
  SOURCE         the document to extract (one file per run)
  --out OUTDIR   output directory, created if missing; use one per source
  --id ID        source id recorded in the manifest (e.g. S01)
  --max-pages N  PDFs only: extract pages 1..N (default: every page)
  --documentation
                 required for a delimited file (.csv, .tsv, or a .txt laid out as
                 a table) and for a JSON file (.json, .jsonl, .ndjson): such a file
                 is as often a dataset as a data dictionary (a JSON file can be an
                 array of records, or a JSON Schema or REDCap metadata export), and
                 nothing in it tells the two apart. Pass it only when the user has
                 said the file describes variables (one row or entry per variable or
                 code), never for a file of records.

Output
  OUTDIR/manifest.json lists every file written and what it holds, plus
  warnings and format-specific stats. stdout carries the same manifest as ONE
  JSON object, except that any list longer than 50 entries is cut to its
  first 50 with a sibling count (warnings_total, ...); files_total is always given.
  Extracted text never goes to stdout; progress and library chatter go to
  stderr. Re-running rewrites the same files: the files listed in OUTDIR's
  previous manifest are removed first, and nothing else in OUTDIR is ever
  deleted or overwritten.

  kind      files                                 locator (cite as)
  xlsx/xls  sheet-NN-<slug>.csv (one per sheet),  _row = Excel row; columns are the
            sheets.json, comments.csv if any      Excel letters: "<sheet> row N"
  csv/tsv   table.csv                             _row = source line (header = 1)
  txt/md    text.md, plus table.csv when a .txt   "L<N>" = source line N
            is consistently delimited
  pdf       pages/pNNNN.txt, text.md              "p.N" = physical page N
  docx      text.md, table-NN.csv, notes.md if    "¶N" = Nth body paragraph,
            footnotes/headers/footers exist       "table NN row R"
  html      text.md                               "§<heading>", or "L<N>" of text.md
  rtf       text.md                               "L<N>" of the converted text
  code      text.md (.sas .do .r .sps .py .sql;   "L<N>" = source line N
            embedded data lines blanked)
  json/xml  pretty.json / pretty.xml              "L<N>" of the pretty copy
  jsonl     text.md (.jsonl .ndjson, as text)     "L<N>" = source line N

  A line break inside one cell or paragraph is written as the two characters
  \\n, so each physical line is one row, paragraph or text line. Workbooks:
  values, not formulas (a formula Excel never calculated reads as empty);
  merged values repeated in every cell of the range; dates as ISO 8601; the
  header row is NOT guessed in the CSV, see stats.sheets[].header_guess.
  PDF pages with under 20 characters of text are listed in
  stats.empty_pages: likely scanned images, view them directly.

Refused (exit 2)
  Data files (.sav .zsav .por .dta .sas7bdat .xpt .rds .rdata .rda .parquet
  .feather .arrow .db .sqlite .sqlite3 .h5 .hdf5, also recognised by content),
  delimited and JSON files (.json .jsonl .ndjson) without --documentation,
  unknown formats, files over 200 MB, password-protected files.

Exit codes
  0  extracted: stdout is the manifest, {"ok": true, ...}
  1  usage or input error: bad arguments, missing/unreadable/empty/corrupt
     source, missing dependency, unsafe --out (inside the schemify skill, or
     holding files this script did not write)
  2  unsupported or refused file: data file, unknown format, over 200 MB,
     password-protected
  Every failure prints {"ok": false, "error": "...", "hint": "<next move>"}.
"""

import argparse
import codecs
import contextlib
import csv
import datetime as dt
import hashlib
import importlib
import io
import json
import logging
import os
import re
import shutil
import subprocess
import sys
import unicodedata
import warnings
import zipfile
from collections import Counter
from html.parser import HTMLParser
from pathlib import Path

MAX_BYTES = 200 * 1024 * 1024
XML_PRETTY_MAX = 50 * 1024 * 1024
# Workbooks whose sheets hold more XML than this are streamed (no merges/hidden info):
# a full load peaks at roughly 12-20x the uncompressed sheet XML in memory.
XLSX_STREAM_XML = 64 * 1024 * 1024
STDOUT_CAP = 50
OUTLINE_CAP = 200
EMPTY_PAGE_CHARS = 20
MERGE_FILL_MAX = 100_000
PIP_FALLBACK = "python3 -m pip install --user openpyxl pypdf python-docx xlrd striprtf"
SUPPORTED = (".xlsx .xlsm .xls .csv .tsv .txt .md .pdf .docx .html .htm .rtf .json .jsonl .ndjson .xml "
             ".sas .do .ado .r .rmd .qmd .sps .py .sql")

EXT_KIND = {
    ".xlsx": "xlsx", ".xlsm": "xlsx", ".xltx": "xlsx", ".xltm": "xlsx", ".xls": "xls",
    ".csv": "csv", ".tsv": "tsv", ".tab": "tsv", ".pdf": "pdf", ".docx": "docx",
    ".html": "html", ".htm": "html", ".xhtml": "html", ".rtf": "rtf",
    ".txt": "txt", ".text": "txt", ".md": "md", ".markdown": "md",
    ".json": "json", ".jsonl": "jsonl", ".ndjson": "jsonl", ".xml": "xml",
    ".sas": "code", ".do": "code", ".ado": "code", ".r": "code", ".rmd": "code", ".qmd": "code",
    ".sps": "code", ".py": "code", ".sql": "code",
}
DATA_EXTS = {".sav", ".zsav", ".por", ".dta", ".sas7bdat", ".xpt", ".rds", ".rdata", ".rda",
             ".parquet", ".feather", ".arrow", ".db", ".sqlite", ".sqlite3", ".h5", ".hdf5"}
DATA_MAGIC = [  # (offset, signature, format) - catches renamed data files
    (0, b"SQLite format 3\x00", "SQLite"), (0, b"PAR1", "Parquet"), (0, b"$FL2", "SPSS"),
    (0, b"$FL3", "SPSS"), (0, b"\x89HDF\r\n\x1a\n", "HDF5"), (0, b"ARROW1", "Arrow"),
    (0, b"FEA1", "Feather"), (0, b"<stata_dta>", "Stata"), (0, b"RDX2\n", "R data"),
    (0, b"RDX3\n", "R data"), (0, b"HEADER RECORD*******LIB", "SAS transport"),
    (12, b"\xc2\xea\x81\x60\xb3\x14\x11\xcf\xbd\x92\x08\x00\x09\xc7\x31\x8c\x18\x1f\x10\x11",
     "SAS7BDAT"),
]
DATA_ERROR = ("This looks like a data file, not a metadata document. "
              "The agent must not open data files.")
DATA_HINT = ("If it embeds variable/value labels you need, ask the user to export a "
             "metadata-only codebook (variable names, labels, value labels — no values or "
             "frequencies) and extract that instead.")
SPECIAL_HINTS = {
    ".doc": "Legacy Word .doc is not supported: save it as .docx or PDF (Word: File > Save As; "
            "macOS: `textutil -convert docx FILE.doc`), then extract that copy.",
    ".xlsb": "Save the workbook as .xlsx (Excel: File > Save As), then extract that copy.",
    ".ods": "Save the spreadsheet as .xlsx (File > Save As), then extract that copy.",
    ".numbers": "Export to Excel (Numbers: File > Export To > Excel), then extract the .xlsx.",
    ".odt": "Save the document as .docx or PDF, then extract that copy.",
    ".pages": "Export to Word or PDF (Pages: File > Export To), then extract that copy.",
    ".ppt": "Save the slides as PDF, then extract the PDF.",
    ".pptx": "Save the slides as PDF, then extract the PDF.",
    ".key": "Export the slides as PDF, then extract the PDF.",
}
ARCHIVE_HINT = "Unpack the archive, then extract each document in it separately."
for _ext in (".zip", ".gz", ".tgz", ".bz2", ".xz", ".7z", ".rar", ".tar"):
    SPECIAL_HINTS[_ext] = ARCHIVE_HINT
# Only paths of these shapes are ever deleted when a previous run's outputs are cleared.
GENERATED = re.compile(
    r"^(?:sheet-\d{2,}-[a-z0-9-]+\.csv|sheets\.json|comments\.csv|table\.csv|table-\d{2,}\.csv"
    r"|text\.md|notes\.md|pretty\.json|pretty\.xml|source\.xml|pages/p\d{4,}\.txt)$")
LOCATORS = {
    "sheets": "Each sheet-NN-<slug>.csv is the NNth sheet: column _row is the Excel row number "
              "and columns A, B, C… are the Excel column letters, so cite '<sheet> row N' or "
              "'<sheet>!C41' (line breaks inside a cell appear as \\n).",
    "pdf": "pages/pNNNN.txt holds physical PDF page N, which text.md marks with a "
           "<!-- page N --> line; cite as p.N (the physical page, which can differ from the "
           "printed page number).",
    "docx": "text.md lines start with ¶N, the paragraph's 1-based position among all body "
            "paragraphs (empty ones count but are not shown); a [table-NN.csv: …] line marks "
            "where table NN sits and its _row is the table row, so cite '¶N' or "
            "'table NN row R'.",
    "html": "text.md line N is the Nth line of the page's text in reading order (headings "
            "as #, list items as -, one table row per line with cells joined by ' | '); cite "
            "as §<nearest heading>, adding L<N> when the section is long.",
    "rtf": "text.md line N is line N of the RTF converted to plain text; cite as L<N>.",
    "plain": "text.md line N is line N of the source; cite as L<N>.",
    "json": "pretty.json line N is line N of the pretty-printed JSON; cite as L<N> or by the "
            "JSON path.",
    "xml": "{file} line N is line N of the XML copy; cite as L<N> or by the element path.",
}

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
MC_ALT = "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"
MC_CHOICE = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Choice"
MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"


def w(tag):
    return "{%s}%s" % (W_NS, tag)


# Parts of a paragraph that hold no visible text (or deleted text) - never walked.
DOCX_SKIP = {w("pPr"), w("rPr"), w("del"), w("moveFrom"), w("instrText"), w("delText"),
             w("sdtPr"), w("sdtEndPr"), w("fldData"), MC_FALLBACK}


class Fail(Exception):
    """A clean failure: printed as {"ok": false, "error", "hint"} with exit code `code`."""

    def __init__(self, error, hint, code=1):
        super().__init__(error)
        self.error, self.hint, self.code = error, hint, code


class _Parser(argparse.ArgumentParser):
    def error(self, message):
        raise Fail("Invalid arguments: " + message,
                   "Usage: extract.py SOURCE --out OUTDIR [--id ID] [--max-pages N] "
                   "[--documentation]; see --help.", 1)


def build_parser():
    p = _Parser(prog="extract.py", description=__doc__,
                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("source", metavar="SOURCE", help="the metadata document to extract")
    p.add_argument("--out", required=True, metavar="OUTDIR",
                   help="output directory (created if missing; one per source)")
    p.add_argument("--id", metavar="ID", help="source id to record in the manifest, e.g. S01")
    p.add_argument("--max-pages", type=int, metavar="N",
                   help="PDFs only: extract at most the first N pages")
    p.add_argument("--include-sheet", action="append", default=[], metavar="NAME",
                   help="workbooks: extract this sheet even though it looks like records "
                        "(only after the user says it describes variables); repeatable")
    p.add_argument("--documentation", action="store_true",
                   help="the user confirmed this file describes variables, not records "
                        "(required for .csv, .tsv, tabular .txt, .json, .jsonl and .ndjson)")
    return p


# ----------------------------------------------------------------------------- helpers

def need(module, kind):
    try:
        return importlib.import_module(module)
    except ImportError as e:
        raise Fail(f"A Python package needed for {kind} files is missing ({e}).",
                   f"Run via `uv run extract.py …` (uv installs the dependencies), or install "
                   f"them with `{PIP_FALLBACK}` and re-run with `python3 extract.py …`.", 1)


def version_of(module):
    return getattr(module, "__version__", None) or "?"


def one_line(s):
    """Escape line breaks so a cell or paragraph stays on one physical line."""
    return s.replace("\r\n", "\n").replace("\r", "\n").replace("\n", "\\n")


def norm_newlines(s):
    return s.replace("\r\n", "\n").replace("\r", "\n")


def collapse(s):
    return re.sub(r"\s+", " ", s).strip()


def plural(n, word):
    return f"{n} {word}{'' if n == 1 else 's'}"


def col_letter(n):
    s = ""
    while n > 0:
        n, rem = divmod(n - 1, 26)
        s = chr(65 + rem) + s
    return s


def col_index(letters):
    n = 0
    for ch in letters.upper():
        n = n * 26 + ord(ch) - 64
    return n


def ranges(nums, fmt=str):
    """[1, 2, 3, 7] -> ['1-3', '7']"""
    nums, out, i = sorted(set(nums)), [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        out.append(fmt(nums[i]) if i == j else f"{fmt(nums[i])}-{fmt(nums[j])}")
        i = j + 1
    return out


def brief(items, cap=20):
    text = ", ".join(items[:cap])
    return text + (f", … ({len(items) - cap} more)" if len(items) > cap else "")


def subject(word, spans):
    """('row', ['5']) -> 'row 5 is'; ('row', ['5-9']) -> 'rows 5-9 are'."""
    if len(spans) == 1 and "-" not in spans[0]:
        return f"{word} {spans[0]} is"
    return f"{word}s {brief(spans)} are"


def slug(name, cap=40):
    s = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")[:cap].strip("-")
    return s or "sheet"


def num_text(f):
    if f == f and abs(f) != float("inf") and f.is_integer() and abs(f) < 1e16:
        return str(int(f))
    return repr(f)


def cell_text(v):
    """A spreadsheet value as CSV text: exact numbers, ISO dates, one line."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return num_text(v)
    if isinstance(v, dt.datetime):
        if v.tzinfo is None and v.time() == dt.time(0):
            return v.date().isoformat()
        return v.isoformat()
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    if isinstance(v, dt.timedelta):
        return str(v)
    s = str(v)
    return one_line(s) if s.strip() else ""


def decode_text(data):
    """Bytes -> (text, encoding): BOMs, BOM-less UTF-16, then utf-8, cp1252, latin-1."""
    for bom, enc in ((codecs.BOM_UTF32_LE, "utf-32"), (codecs.BOM_UTF32_BE, "utf-32"),
                     (codecs.BOM_UTF16_LE, "utf-16"), (codecs.BOM_UTF16_BE, "utf-16")):
        if data.startswith(bom):
            try:
                return data.decode(enc), enc
            except UnicodeDecodeError:
                break
    sample = data[:4096]
    if len(sample) >= 4 and sample.count(0) > len(sample) // 4:
        enc = "utf-16-le" if sample[1::2].count(0) > sample[0::2].count(0) else "utf-16-be"
        try:
            return data.decode(enc), enc
        except UnicodeDecodeError:
            pass
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
        except UnicodeDecodeError:
            continue
        if enc == "utf-8-sig" and not data.startswith(codecs.BOM_UTF8):
            enc = "utf-8"
        return text, enc
    raise AssertionError("latin-1 decodes any bytes")


def now_iso():
    return dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat().replace(
        "+00:00", "Z")


def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


class Run:
    """Collects the files, warnings and stats of one extraction."""

    def __init__(self, out, size):
        self.out, self.size = out, size
        self.files, self.warnings, self.stats = [], [], {}

    def new(self, rel, what=""):
        path = self.out / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        if self.out.resolve() not in (path.parent.resolve(), *path.parent.resolve().parents):
            raise Fail(f"{path.parent} leads outside {self.out} (a symlink?); refusing to write "
                       f"there.", "Use a new, empty --out directory.", 1)
        if path.exists() or path.is_symlink():
            raise Fail(f"{path} already exists and no earlier extract.py manifest lists it; "
                       f"refusing to overwrite it.",
                       "Use a new or dedicated --out directory (one per source document), or "
                       "move that file out of the way and re-run.", 1)
        entry = {"path": rel, "what": what}
        self.files.append(entry)  # listed before writing, so a crash can still be cleaned up
        return path, entry

    def text(self, rel, text, what):
        path, entry = self.new(rel, what)
        with open(path, "w", encoding="utf-8", newline="\n") as f:
            f.write(text)
        return entry

    @contextlib.contextmanager
    def csv(self, rel):
        path, entry = self.new(rel)
        with open(path, "w", encoding="utf-8", newline="") as f:
            yield csv.writer(f, lineterminator="\n"), entry

    def warn(self, message):
        self.warnings.append(message)


# ------------------------------------------------------------------ safety and routing

NAME_LINE = re.compile(r"""^name:\s*["']?([^"'#\s]+)["']?\s*(?:#.*)?$""")


def is_schemify_skill(skill_md):
    try:
        if not skill_md.is_file():
            return False
        with open(skill_md, encoding="utf-8", errors="replace") as f:
            lines = f.read(8192).splitlines()
    except OSError:
        return False
    if not lines or lines[0].strip() != "---":
        return False
    for line in lines[1:]:
        if line.strip() == "---":
            break
        m = NAME_LINE.match(line.strip())
        if m:
            return m.group(1) == "schemify"
    return False


def schemify_skill_above(out):
    seen = set()
    for base in (Path(os.path.abspath(out)), out.resolve()):
        for d in (base, *base.parents):
            if d not in seen:
                seen.add(d)
                if is_schemify_skill(d / "SKILL.md"):
                    return d
    return None


def previous_manifest(out):
    mf = out / "manifest.json"
    if not mf.exists():
        return None
    try:
        m = json.loads(mf.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        m = None
    if not (isinstance(m, dict) and isinstance(m.get("files"), list)
            and {"kind", "sha256", "extracted_at"} <= m.keys()):
        raise Fail(f"{mf} exists but was not written by extract.py; refusing to touch that "
                   f"directory.",
                   "Pass a new or dedicated --out directory (one per source document).", 1)
    return m


def clear_previous(out, prev):
    """Delete the files a previous run listed (only paths shaped like our outputs)."""
    for entry in prev.get("files", []) if prev else []:
        rel = entry.get("path") if isinstance(entry, dict) else None
        if isinstance(rel, str) and GENERATED.match(rel):
            p = out / rel
            if p.is_symlink() or p.is_file():
                p.unlink()
    with contextlib.suppress(OSError):
        (out / "pages").rmdir()  # only succeeds when empty
    with contextlib.suppress(OSError):
        (out / "manifest.json").unlink()


def data_signature(head):
    for offset, magic, label in DATA_MAGIC:
        if head[offset:offset + len(magic)] == magic:
            return label
    return None


def head_text(head):
    for bom, enc in ((codecs.BOM_UTF8, "utf-8"), (codecs.BOM_UTF16_LE, "utf-16-le"),
                     (codecs.BOM_UTF16_BE, "utf-16-be")):
        if head.startswith(bom):
            return head[len(bom):].decode(enc, "ignore")
    return head.decode("latin-1")


def sniff(path, head):
    """What the bytes say the file is, whatever its extension."""
    if head.lstrip()[:5] == b"%PDF-":
        return "pdf"
    if head.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(path) as z:
                names = set(z.namelist())
        except (zipfile.BadZipFile, OSError):
            return "zip"
        if "xl/workbook.xml" in names:
            return "xlsx"
        if "xl/workbook.bin" in names:
            return "xlsb"
        return "docx" if "word/document.xml" in names else "zip"
    if head.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return "ole2"
    text = head_text(head).lstrip()
    if text.startswith("{\\rtf"):
        return "rtf"
    low = text[:4096].lower()
    if low.startswith("<"):
        if "<html" in low or "<!doctype html" in low:
            return "html"
        if low.startswith("<?xml") or re.match(r"<[a-z_]", low):
            return "xml"
    return None


def detect_kind(path, ext, head):
    """Route by extension, corrected by content. Returns (kind, note-or-None)."""
    kind = EXT_KIND.get(ext)
    found = sniff(path, head)
    shown = ext or "(none)"
    if found == "xlsb":
        raise Fail("Binary Excel workbooks (.xlsb) are not supported.",
                   f"{SPECIAL_HINTS['.xlsb']} Supported: {SUPPORTED}.", 2)
    if found == "ole2":
        if kind == "xls":
            return kind, None
        if kind in ("xlsx", "csv", "tsv", "txt"):
            return "xls", (f"the content is a legacy Excel (OLE2) file although the extension "
                           f"is {shown}; extracted as xls")
        if kind == "docx" or ext == ".doc":
            raise Fail("This is a legacy Word .doc file (or a password-protected .docx).",
                       f"{SPECIAL_HINTS['.doc']} If it is password-protected, ask the steward "
                       f"for an unprotected copy. Supported: {SUPPORTED}.", 2)
    if found == "zip" and kind is None:
        raise Fail(f"Unsupported file type: {shown} (a zip archive).",
                   f"{ARCHIVE_HINT} Supported: {SUPPORTED}.", 2)
    if found in ("pdf", "xlsx", "docx", "rtf") and found != kind:
        return found, (f"the content is {found} although the extension is {shown}; "
                       f"extracted as {found}")
    if found in ("html", "xml") and found != kind and kind in (None, "xls", "xlsx", "docx"):
        return found, (f"the content is {found} although the extension is {shown}; "
                       f"extracted as {found}")
    if kind is None:
        hint = SPECIAL_HINTS.get(ext) or (
            "Convert the document to a supported format (PDF, .docx or .xlsx usually work) and "
            "extract that copy. If it is a plain-text metadata file (e.g. SAS, SPSS or Stata "
            "syntax that defines labels), save a copy with a .txt extension and extract that; "
            "never do this for a data file.")
        raise Fail(f"Unsupported file type: {shown}.", f"{hint} Supported: {SUPPORTED}.", 2)
    return kind, None


def encrypted_ooxml(path):
    """Password-protected .xlsx/.docx files are OLE2 containers with this stream."""
    try:
        return "EncryptionInfo".encode("utf-16-le") in path.read_bytes()
    except OSError:
        return False


@contextlib.contextmanager
def quiet_libraries(run):
    """Route library warnings/log records into stats instead of the terminal."""
    messages = []

    class Collect(logging.Handler):
        def emit(self, record):
            messages.append(record.getMessage())

    handler = Collect(level=logging.WARNING)
    loggers = [logging.getLogger(n) for n in ("pypdf", "openpyxl", "docx")]
    saved = [lg.propagate for lg in loggers]
    for lg in loggers:
        lg.addHandler(handler)
        lg.propagate = False
    try:
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            try:
                yield messages
            finally:
                messages.extend(str(c.message) for c in caught)
    finally:
        for lg, prop in zip(loggers, saved):
            lg.removeHandler(handler)
            lg.propagate = prop
        earlier = run.stats.get("library_warnings", [])
        unique = list(dict.fromkeys(m.strip() for m in earlier + messages if m and m.strip()))
        if unique:
            run.stats["library_warnings"] = unique[:10]
            if len(unique) > 10:
                run.stats["library_warnings_total"] = len(unique)


# ------------------------------------------------------------------------ workbooks

DICT_WORDS = {"variable", "variables", "var", "varname", "name", "label", "labels", "description", "desc",
              "type", "datatype", "format", "value", "values", "code", "codes", "coding", "min", "max", "minimum",
              "maximum", "range", "unit", "units", "note", "notes", "comment", "comments", "question", "item",
              "section", "category", "domain", "form", "instrument", "source", "missing", "length", "width",
              "decimals", "level", "levels", "meaning", "definition", "universe", "skip", "routing", "field"}
IDENT_RE = re.compile(r"^[A-Za-z][A-Za-z0-9_.]*$")
NUMLIKE_RE = re.compile(r"^[-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?$|^\d{4}-\d{2}-\d{2}")


def header_row_of(raw):
    """The first of a sheet's first 30 non-empty rows that is at least 60% as wide as the widest."""
    counts = [(r, len(raw[r])) for r in sorted(raw)[:30]]
    if not counts:
        return None
    top = max(n for _, n in counts)
    return next(r for r, n in counts if n >= 0.6 * top)


def body_names(raw):
    """Identifier-like cells below a sheet's heading row: the variable names a dictionary sheet lists."""
    head = header_row_of(raw)
    return {v.lower() for r, cols in raw.items() if head is None or r > head
            for v in cols.values() if IDENT_RE.match(v)}


def records_reason(raw, elsewhere=None):
    """Why a sheet looks like records (one row per person or visit) rather than documentation, or None.
    Such a sheet is real data (SKILL.md, law 2): it is not written out unless --include-sheet names it."""
    head = header_row_of(raw)
    if head is None:
        return None
    header = [v for _, v in sorted(raw[head].items())]
    if len(header) < 3:
        return None
    dict_like = sum(1 for h in header if (words := set(re.findall(r"[a-z]+", h.lower()))) and words <= DICT_WORDS)
    if dict_like >= 2:
        return None
    if elsewhere and sum(1 for h in header if h.lower() in elsewhere) >= 0.6 * len(header):
        return "its column headings are variable names that another sheet lists"
    body = sorted(r for r in raw if r > head)
    if len(body) < 5 or sum(1 for h in header if IDENT_RE.match(h)) < 0.8 * len(header):
        return None
    cells = [v for r in body[:200] for v in raw[r].values()]
    if cells and sum(1 for v in cells if NUMLIKE_RE.match(v)) >= 0.6 * len(cells):
        return "a heading row of variable names over rows of mostly numbers"
    return None


def skipped_sheet(idx, name, state, reason):
    return {"index": idx, "name": name, "file": None, "state": state, "skipped": reason}


def write_sheet(run, idx, name, state, raw, merged=(), hidden_rows=(), hidden_cols=()):
    """raw: {row: {col: text}} with non-empty cells only (1-based). Writes one CSV."""
    ordered = sorted(raw)
    header_guess = header_row_of(raw)
    grid = raw  # merged values are filled in place; ranges never overlap
    for r1, c1, r2, c2 in merged:
        value = raw.get(r1, {}).get(c1)
        if not value:
            continue
        if (r2 - r1 + 1) * (c2 - c1 + 1) > MERGE_FILL_MAX:
            run.warn(f"sheet '{name}': merged range {col_letter(c1)}{r1}:{col_letter(c2)}{r2} "
                     f"is too large to fill; its value appears only in the top-left cell.")
            continue
        for r in range(r1, r2 + 1):
            row = grid.setdefault(r, {})
            for c in range(c1, c2 + 1):
                row[c] = value
    used = set()
    for cols in grid.values():
        used.update(cols)
    max_col = max(used, default=0)
    rel = f"sheet-{idx:02d}-{slug(name)}.csv"
    with run.csv(rel) as (out, entry):
        out.writerow(["_row"] + [col_letter(c) for c in range(1, max_col + 1)])
        for r in sorted(grid):
            cols = grid[r]
            out.writerow([r] + [cols.get(c, "") for c in range(1, max_col + 1)])
        size = f"{plural(len(grid), 'row')} x {plural(max_col, 'col')}" if grid else "empty"
        entry["what"] = f"sheet '{name}' ({size}{'' if state == 'visible' else ', ' + state})"
    rows = sorted(grid)
    return {
        "index": idx, "name": name, "file": rel, "state": state,
        "dims": (f"{col_letter(min(used))}{rows[0]}:{col_letter(max_col)}{rows[-1]}"
                 if grid else None),
        "rows": len(grid), "cols": max_col,
        "first_nonempty_row": ordered[0] if ordered else None,
        "header_guess": header_guess,
        "merged": [f"{col_letter(c1)}{r1}:{col_letter(c2)}{r2}" for r1, c1, r2, c2 in merged],
        "hidden_rows": ranges([r for r in hidden_rows if r in grid]),
        "hidden_cols": ranges([c for c in hidden_cols if c in used], col_letter),
    }


def stream_sheet(run, idx, name, state, ws, include=False):
    """Two passes over a read-only worksheet: shape first, then the CSV. Low memory."""
    ws.reset_dimensions()

    def rows():
        for r, values in enumerate(ws.iter_rows(values_only=True), 1):
            cols = {c: s for c, s in ((c, cell_text(v)) for c, v in enumerate(values, 1)) if s}
            if cols:
                yield r, cols

    sample, n, first, last, min_col, max_col = {}, 0, None, None, None, 0
    for r, cols in rows():
        if len(sample) < 230:
            sample[r] = cols
        n, last, first = n + 1, r, first or r
        min_col = min(min(cols), min_col or min(cols))
        max_col = max(max_col, max(cols))
    header_guess = header_row_of(sample)
    reason = None if include else records_reason(sample)
    if reason:
        return skipped_sheet(idx, name, state, reason)
    rel = f"sheet-{idx:02d}-{slug(name)}.csv"
    with run.csv(rel) as (out, entry):
        out.writerow(["_row"] + [col_letter(c) for c in range(1, max_col + 1)])
        for r, cols in rows():
            out.writerow([r] + [cols.get(c, "") for c in range(1, max_col + 1)])
        size = f"{plural(n, 'row')} x {plural(max_col, 'col')}" if n else "empty"
        entry["what"] = f"sheet '{name}' ({size}{'' if state == 'visible' else ', ' + state})"
    return {"index": idx, "name": name, "file": rel, "state": state,
            "dims": f"{col_letter(min_col)}{first}:{col_letter(max_col)}{last}" if n else None,
            "rows": n, "cols": max_col, "first_nonempty_row": first,
            "header_guess": header_guess, "merged": [], "hidden_rows": [], "hidden_cols": []}


def finish_workbook(run, sheets, comments):
    run.text("sheets.json", json.dumps(sheets, indent=2, ensure_ascii=False) + "\n",
             "per-sheet stats: dims, first_nonempty_row, header_guess, merged ranges, "
             "hidden rows and columns")
    if comments:
        with run.csv("comments.csv") as (out, entry):
            out.writerow(["sheet", "cell", "row", "author", "comment"])
            for sheet, r, c, author, text in comments:
                out.writerow([sheet, f"{col_letter(c)}{r}", r, author, text])
            entry["what"] = f"{plural(len(comments), 'cell comment')} (sheet, cell, row, author)"
    for s in sheets:
        if s.get("skipped"):
            run.warn(f"sheet '{s['name']}' was NOT extracted: it looks like records, not documentation "
                     f"({s['skipped']}). Never read it. Ask the user; only if they say it describes "
                     f"variables, re-run with --include-sheet \"{s['name']}\".")
            continue
        if s["state"] != "visible":
            run.warn(f"sheet '{s['name']}' is {s['state']} in Excel; extracted anyway "
                     f"({s['file']}).")
        if s["hidden_rows"]:
            run.warn(f"sheet '{s['name']}': {subject('row', s['hidden_rows'])} hidden in Excel "
                     f"but has content; extracted anyway.")
        if s["hidden_cols"]:
            run.warn(f"sheet '{s['name']}': {subject('column', s['hidden_cols'])} hidden in "
                     f"Excel but has content; extracted anyway.")
    run.stats["sheets"] = sheets
    run.stats["comments"] = len(comments)
    run.stats["skipped_sheets"] = [s["name"] for s in sheets if s.get("skipped")]


def sheet_xml_bytes(path):
    try:
        with zipfile.ZipFile(path) as z:
            return sum(i.file_size for i in z.infolist()
                       if i.filename.startswith("xl/worksheets/") and i.filename.endswith(".xml"))
    except (zipfile.BadZipFile, OSError):
        return 0


def extract_xlsx(src, run, args):
    openpyxl = need("openpyxl", "xlsx")
    stream = sheet_xml_bytes(src) > XLSX_STREAM_XML
    try:
        wb = openpyxl.load_workbook(src, read_only=stream, data_only=True, keep_links=False)
    except Exception as e:
        raise Fail(f"Could not open {src.name} as an Excel workbook ({type(e).__name__}: {e}).",
                   "Check that it opens in Excel; re-save it as .xlsx (or save each sheet as "
                   "CSV) and extract that copy.", 1)
    run.stats["engine"] = f"openpyxl {version_of(openpyxl)}"
    run.stats["mode"] = "streamed" if stream else "full"
    if stream:
        run.warn(f"the workbook's sheets hold over {XLSX_STREAM_XML >> 20} MB of XML, so it was "
                 f"streamed: merged ranges, hidden rows/columns and cell comments were not "
                 f"read, and merged values appear only in each range's top-left cell.")
    sheets, comments, pending = [], [], []
    include = set(args.include_sheet or [])
    try:
        for idx, name in enumerate(wb.sheetnames, 1):
            ws = wb[name]
            if not hasattr(ws, "iter_rows"):
                run.warn(f"sheet {idx} '{name}' is a chart sheet; it has no cells to extract.")
                continue
            state = getattr(ws, "sheet_state", "visible") or "visible"
            if stream:
                sheets.append(stream_sheet(run, idx, name, state, ws, include=name in include))
                continue
            raw = {}
            cells = getattr(ws, "_cells", None)  # existing cells only; iter_rows would fill gaps
            pairs = (cells.items() if isinstance(cells, dict)
                     else (((c.row, c.column), c) for row in ws.iter_rows() for c in row))
            for (r, c), cell in pairs:
                text = cell_text(cell.value)
                if text:
                    raw.setdefault(r, {})[c] = text
                note = getattr(cell, "comment", None)
                if note is not None and (note.text or "").strip():
                    comments.append((name, r, c, note.author or "", one_line(note.text.strip())))
            merged = [(m.min_row, m.min_col, m.max_row, m.max_col)
                      for m in ws.merged_cells.ranges]
            hidden_rows = [r for r, d in ws.row_dimensions.items() if d.hidden]
            hidden_cols = set()
            for key, d in ws.column_dimensions.items():
                if d.hidden:
                    lo = d.min or col_index(key)
                    hidden_cols.update(range(lo, min(d.max or lo, 16384) + 1))
            pending.append((idx, name, state, raw, merged, hidden_rows, hidden_cols))
    finally:
        if stream:
            wb.close()
    sheets += write_sheets(run, pending, include)
    skipped = {x["name"] for x in sheets if x.get("skipped")}
    comments = [c for c in comments if c[0] not in skipped]
    comments.sort(key=lambda c: (wb.sheetnames.index(c[0]), c[1], c[2]))
    finish_workbook(run, sheets, comments)
    return LOCATORS["sheets"]


def write_sheets(run, pending, include):
    """Write each collected sheet, or skip it when it looks like records: by its own shape, or because
    its headings are variable names another sheet of the workbook lists."""
    names = [body_names(p[3]) for p in pending]
    out = []
    for i, (idx, name, state, raw, merged, hidden_rows, hidden_cols) in enumerate(pending):
        elsewhere = set().union(*(n for j, n in enumerate(names) if j != i))
        reason = None if name in include else records_reason(raw, elsewhere)
        out.append(skipped_sheet(idx, name, state, reason) if reason else
                   write_sheet(run, idx, name, state, raw, merged, hidden_rows, hidden_cols))
    return out


def xls_text(xlrd, ctype, value, datemode):
    if ctype == xlrd.XL_CELL_TEXT:
        return one_line(value) if value.strip() else ""
    if ctype == xlrd.XL_CELL_NUMBER:
        return num_text(value)
    if ctype == xlrd.XL_CELL_DATE:
        try:
            when = xlrd.xldate.xldate_as_datetime(value, datemode)
        except Exception:
            return num_text(value)
        return when.time().isoformat() if value < 1 else cell_text(when)
    if ctype == xlrd.XL_CELL_BOOLEAN:
        return "TRUE" if value else "FALSE"
    if ctype == xlrd.XL_CELL_ERROR:
        return getattr(xlrd, "error_text_from_code", {}).get(value, "#ERROR")
    return ""


def extract_xls(src, run, args):
    xlrd = need("xlrd", "xls")
    log = io.StringIO()  # xlrd logs to stdout by default, which would corrupt the JSON
    formatting = True
    try:
        book = xlrd.open_workbook(str(src), formatting_info=True, on_demand=True, logfile=log)
    except Exception:
        formatting = False
        try:
            book = xlrd.open_workbook(str(src), on_demand=True, logfile=log)
        except Exception as e:
            if "encrypt" in str(e).lower() or encrypted_ooxml(src):
                raise Fail("This workbook is password-protected.",
                           "Ask the steward for an unprotected copy (Excel: File > Info > "
                           "Protect Workbook > Encrypt with Password, cleared), then extract "
                           "that.", 2)
            raise Fail(f"Could not open {src.name} as a legacy Excel workbook "
                       f"({type(e).__name__}: {e}).",
                       "Re-save it in Excel as .xlsx and extract that copy.", 1)
        run.warn("merged ranges, hidden rows/columns and comments could not be read from this "
                 ".xls; merged values appear only in each range's top-left cell.")
    run.stats["engine"] = f"xlrd {version_of(xlrd)}"
    sheets, comments, pending = [], [], []
    for i in range(book.nsheets):
        sh = book.sheet_by_index(i)
        raw = {}
        for r in range(sh.nrows):
            for c, (t, v) in enumerate(zip(sh.row_types(r), sh.row_values(r))):
                text = xls_text(xlrd, t, v, book.datemode)
                if text:
                    raw.setdefault(r + 1, {})[c + 1] = text
        merged, hidden_rows, hidden_cols = [], [], []
        if formatting:
            merged = [(rlo + 1, clo + 1, rhi, chi) for rlo, rhi, clo, chi in sh.merged_cells]
            hidden_rows = [r + 1 for r, info in sh.rowinfo_map.items() if info.hidden]
            hidden_cols = [c + 1 for c, info in sh.colinfo_map.items() if info.hidden]
            for (r, c), note in sorted(getattr(sh, "cell_note_map", {}).items()):
                if (note.text or "").strip():
                    comments.append((sh.name, r + 1, c + 1, note.author or "",
                                     one_line(note.text.strip())))
        state = {1: "hidden", 2: "veryHidden"}.get(sh.visibility, "visible")
        pending.append((i + 1, sh.name, state, raw, merged, hidden_rows, hidden_cols))
        book.unload_sheet(i)
    book.release_resources()
    sheets = write_sheets(run, pending, set(args.include_sheet or []))
    skipped = {x["name"] for x in sheets if x.get("skipped")}
    comments = [c for c in comments if c[0] not in skipped]
    if log.getvalue().strip():
        run.stats["library_warnings"] = log.getvalue().strip().splitlines()[:10]
    finish_workbook(run, sheets, comments)
    return LOCATORS["sheets"]


# -------------------------------------------------------------------- delimited text

def write_table(run, text, delimiter):
    """Write table.csv (UTF-8, comma) with _row = the source line each record starts on."""
    csv.field_size_limit(16 * 1024 * 1024)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    header, header_line, rows, widest, ragged, runaway, prev = None, None, 0, 0, [], [], 0
    with run.csv("table.csv") as (out, entry):
        try:
            for record in reader:
                start, prev = prev + 1, reader.line_num
                if not any(field.strip() for field in record):
                    continue
                if prev - start > 50:
                    runaway.append(start)
                record = [one_line(field) for field in record]
                widest = max(widest, len(record))
                if header is None:
                    header, header_line = record, start
                    out.writerow(["_row"] + record)
                    continue
                if len(record) != len(header):
                    ragged.append(start)
                out.writerow([start] + record)
                rows += 1
        except csv.Error as e:
            raise Fail(f"Could not parse the file as delimited text near line "
                       f"{reader.line_num} ({e}).",
                       "Open it in a spreadsheet application and re-save it as CSV (or .xlsx), "
                       "then extract that copy.", 1)
        entry["what"] = (f"{plural(rows, 'row')} x {plural(len(header or []), 'col')} below the "
                         f"header (line {header_line}); UTF-8, comma-separated")
    if ragged:
        run.warn(f"{plural(len(ragged), 'row')} {'has' if len(ragged) == 1 else 'have'} a "
                 f"different number of fields than the header (first at line {ragged[0]}); "
                 f"check for unquoted delimiters there.")
    if runaway:
        run.warn(f"records starting at lines {brief(ranges(runaway))} span more than 50 lines, "
                 f"which usually means an unbalanced quote; check the source there.")
    run.stats.update({"delimiter": delimiter, "header_line": header_line, "rows": rows,
                      "columns": len(header or []), "max_fields": widest,
                      "ragged_rows": len(ragged)})
    return header_line


def sniff_delimiter(text, prefer=None):
    sample = text[:65536]
    if len(text) > 65536 and "\n" in sample:
        sample = sample[:sample.rfind("\n")]
    if prefer and prefer in sample:
        return prefer
    try:
        found = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        if found in ",;\t|":
            return found
    except csv.Error:
        pass
    return prefer or ","


def consistent_delimiter(text):
    """A .txt counts as a table only if one delimiter gives a steady field count."""
    lines = [ln for ln in text.split("\n", 400)[:400] if ln.strip()][:200]
    if len(lines) < 5:
        return None
    best = None
    for d in ("\t", "|", ";", ","):
        counts = Counter()
        for ln in lines:
            try:
                counts[len(next(csv.reader([ln], delimiter=d)))] += 1
            except (csv.Error, StopIteration):
                counts[1] += 1
        width, freq = counts.most_common(1)[0]
        if width >= 2 and freq >= 0.8 * len(lines) and (best is None or freq > best[0]):
            best = (freq, d)
    return best[1] if best else None


def extract_delimited(src, run, args, kind):
    text, encoding = decode_text(src.read_bytes())
    if "\x00" in text:
        text = text.replace("\x00", "")
        run.warn("removed NUL bytes from the text; check the encoding if the output looks wrong.")
    delimiter = sniff_delimiter(text, "\t" if kind == "tsv" else None)
    run.stats["encoding"] = encoding
    header_line = write_table(run, text, delimiter)
    return (f"table.csv column _row is the 1-based line number in the source file (the header "
            f"is line {header_line or 1}); cite as '<file> row N' (line breaks inside a field "
            f"appear as \\n).")


def extract_plain(src, run, args, table_ok=False):
    text, encoding = decode_text(src.read_bytes())
    if "\x00" in text:
        text = text.replace("\x00", "")
        run.warn("removed NUL bytes from the text; check the encoding if the output looks wrong.")
    text = norm_newlines(text)
    if text and not text.endswith("\n"):
        text += "\n"
    lines = text.count("\n")
    run.text("text.md", text, f"the source text, {plural(lines, 'line')} (line N = source line N)")
    run.stats.update({"encoding": encoding, "lines": lines, "chars": len(text)})
    if not text.strip():
        run.warn("the document contains no text.")
    locators = LOCATORS["plain"]
    if table_ok:
        delimiter = consistent_delimiter(text)
        if delimiter:
            write_table(run, text, delimiter)
            locators += " table.csv column _row is that same line number."
    return locators


EMBEDDED_SAS = re.compile(r"^\s*(datalines|cards|lines|parmcards)(4?)\s*;\s*$", re.I)
EMBEDDED_STATA = re.compile(r"^\s*input\b", re.I)


def extract_code(src, run, args):
    """Cleaning or processing code as plain text, line N = source line N. Rows of data embedded in the
    code (SAS datalines/cards, Stata input...end) are replaced line for line: they may be records."""
    text, encoding = decode_text(src.read_bytes())
    lines = norm_newlines(text.replace("\x00", "")).split("\n")
    stata = src.suffix.lower() in (".do", ".ado")
    out, hidden, i = [], [], 0
    while i < len(lines):
        out.append(lines[i])
        m = EMBEDDED_SAS.match(lines[i])
        if not (m or (stata and EMBEDDED_STATA.match(lines[i]))):
            i += 1
            continue
        j = i + 1
        while j < len(lines):
            stripped = lines[j].strip()
            if (stripped.lower() == "end") if not m else (stripped == ";;;;" if m.group(2) else stripped.startswith(";")):
                break
            out.append("[embedded data line, not extracted]")
            j += 1
        if j > i + 1:
            hidden.append((i + 2, j))
        i = j
    text = "\n".join(out)
    if text and not text.endswith("\n"):
        text += "\n"
    n = text.count("\n")
    run.text("text.md", text, f"the code, {plural(n, 'line')} (line N = source line N)")
    run.stats.update({"encoding": encoding, "lines": n, "chars": len(text),
                      "embedded_data_lines": sum(b - a + 1 for a, b in hidden)})
    for a, b in hidden:
        run.warn(f"L{a}-L{b}: rows of data embedded in the code were not extracted (they may be records, "
                 f"SKILL.md law 2). If the user says they are a lookup table, ask for it as its own document.")
    return LOCATORS["plain"]


# ---------------------------------------------------------------------------- PDF

LIGATURES = str.maketrans({"\ufb00": "ff", "\ufb01": "fi", "\ufb02": "fl", "\ufb03": "ffi",
                           "\ufb04": "ffl", "\ufb05": "st", "\ufb06": "st"})


def pdf_outline(reader, run):
    items, total = [], 0

    def walk(nodes, level):
        nonlocal total
        for node in nodes:
            if isinstance(node, list):
                walk(node, level + 1)
                continue
            total += 1
            if len(items) >= OUTLINE_CAP:
                continue
            try:
                page = reader.get_destination_page_number(node)
            except Exception:
                page = None
            items.append({"title": collapse(str(getattr(node, "title", "") or ""))[:200],
                          "page": page + 1 if isinstance(page, int) and page >= 0 else None,
                          "level": level})

    try:
        walk(reader.outline, 1)
    except Exception as e:
        run.warn(f"the PDF outline (bookmarks) could not be read ({type(e).__name__}).")
    return items, total


def extract_pdf(src, run, args):
    pypdf = need("pypdf", "pdf")
    try:
        reader = pypdf.PdfReader(str(src), strict=False)
        encrypted = reader.is_encrypted
    except Exception as e:
        raise Fail(f"Could not open {src.name} as a PDF ({type(e).__name__}: {e}).",
                   "Check that it opens in a PDF viewer; if it does, print it to a new PDF and "
                   "extract that copy.", 1)
    if encrypted:
        try:
            unlocked = reader.decrypt("")
        except Exception as e:
            if "cryptography" in str(e).lower() or "AES" in str(e):
                raise Fail("This PDF is AES-encrypted and the `cryptography` package is "
                           "missing.",
                           "Run via `uv run extract.py …`, or `python3 -m pip install --user "
                           "cryptography`, then re-run.", 1)
            unlocked = 0
        if not unlocked:
            raise Fail("This PDF is password-protected.",
                       "Ask the steward for an unprotected copy; this script does not open "
                       "password-protected files.", 2)
    try:
        total = len(reader.pages)
    except Exception as e:
        raise Fail(f"Could not read the pages of {src.name} ({type(e).__name__}: {e}).",
                   "Print the PDF to a new PDF file and extract that copy.", 1)
    limit = min(total, args.max_pages or total)
    run.stats.update({"engine": f"pypdf {version_of(pypdf)}", "encrypted": bool(encrypted),
                      "pages": total, "pages_extracted": limit})
    if limit < total:
        run.warn(f"extracted pages 1-{limit} of {total} (--max-pages {args.max_pages}); "
                 f"re-run without --max-pages for the rest.")
    chunks, empty, failed = [], [], []
    for i in range(limit):
        n = i + 1
        if total > 100 and n % 100 == 0:
            print(f"extract: page {n}/{limit}", file=sys.stderr, flush=True)
        try:
            text = reader.pages[i].extract_text() or ""
        except Exception:
            text = ""
            failed.append(n)
        text = norm_newlines(text).replace("\x00", "").translate(LIGATURES)
        text = "\n".join(line.rstrip() for line in text.split("\n")).strip("\n")
        chars = len(re.sub(r"\s", "", text))
        if chars < EMPTY_PAGE_CHARS:
            empty.append(n)
        run.text(f"pages/p{n:04d}.txt", text + "\n" if text else "",
                 f"page {n} ({plural(chars, 'char')} of text)")
        chunks.append(f"<!-- page {n} -->\n{text}\n" if text else f"<!-- page {n} -->\n")
    run.text("text.md", "\n".join(chunks),
             f"all {plural(limit, 'extracted page')}, each after a <!-- page N --> line")
    if failed:
        run.warn(f"text extraction failed on pages {brief(ranges(failed))}; view these pages "
                 f"directly.")
    if empty:
        run.warn(f"{subject('page', ranges(empty))} under {EMPTY_PAGE_CHARS} characters of "
                 f"text and may be scanned images: view these pages directly; the text layer "
                 f"is empty.")
    run.stats["empty_pages"] = empty
    outline, outline_total = pdf_outline(reader, run)
    run.stats["outline"] = outline
    run.stats["outline_total"] = outline_total
    try:
        labels = list(reader.page_labels)[:limit]
    except Exception:
        labels = []
    if any(label != str(i + 1) for i, label in enumerate(labels)):
        run.stats["page_labels"] = labels
        odd = next(i for i, label in enumerate(labels) if label != str(i + 1))
        run.warn(f"this PDF prints its own page labels (physical p.{odd + 1} is labelled "
                 f"'{labels[odd]}'); p.N in these files is always the physical page, see "
                 f"stats.page_labels.")
    return LOCATORS["pdf"]


# --------------------------------------------------------------------------- DOCX

def iter_blocks(parent):
    """Body-level paragraphs and tables in document order, unwrapping content controls."""
    for child in parent:
        tag = child.tag
        if tag == w("p"):
            yield "p", child
        elif tag == w("tbl"):
            yield "tbl", child
        elif tag == w("sdt"):
            content = child.find(w("sdtContent"))
            if content is not None:
                yield from iter_blocks(content)
        elif tag == w("customXml"):
            yield from iter_blocks(child)
        elif tag == MC_ALT:
            choice = child.find(MC_CHOICE)
            if choice is not None:
                yield from iter_blocks(choice)


def iter_tagged(parent, tag):
    for child in parent:
        if child.tag == tag:
            yield child
        elif child.tag == w("sdt"):
            content = child.find(w("sdtContent"))
            if content is not None:
                yield from iter_tagged(content, tag)
        elif child.tag == w("customXml"):
            yield from iter_tagged(child, tag)


def wval(el, default=None):
    return default if el is None else el.get(w("val"), default)


def to_int(value, default=0):
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


class Docx:
    """Paragraph text, heading levels and list numbers from WordprocessingML."""

    def __init__(self, doc):
        self.styles = {}
        for st in doc.styles.element.findall(w("style")):
            ppr = st.find(w("pPr"))
            numpr = ppr.find(w("numPr")) if ppr is not None else None
            self.styles[st.get(w("styleId"))] = {
                "name": wval(st.find(w("name")), "") or "",
                "based": wval(st.find(w("basedOn"))),
                "outline": wval(ppr.find(w("outlineLvl"))) if ppr is not None else None,
                "numId": wval(numpr.find(w("numId"))) if numpr is not None else None,
                "ilvl": wval(numpr.find(w("ilvl"))) if numpr is not None else None,
            }
        self.links = {rid: rel.target_ref for rid, rel in doc.part.rels.items() if rel.is_external}
        self.abstract, self.nums, self.counters, self.started = {}, {}, {}, set()
        self.numbered = 0
        try:
            numbering = doc.part.numbering_part.element
        except Exception:
            numbering = None
        if numbering is not None:
            for an in numbering.findall(w("abstractNum")):
                self.abstract[an.get(w("abstractNumId"))] = {
                    to_int(lvl.get(w("ilvl"))): (wval(lvl.find(w("numFmt")), "decimal"),
                                                 wval(lvl.find(w("lvlText")), ""),
                                                 to_int(wval(lvl.find(w("start"))), 1))
                    for lvl in an.findall(w("lvl"))}
            for num in numbering.findall(w("num")):
                overrides = {}
                for ov in num.findall(w("lvlOverride")):
                    start = ov.find(w("startOverride"))
                    if start is not None:
                        overrides[to_int(ov.get(w("ilvl")))] = to_int(wval(start), 1)
                self.nums[num.get(w("numId"))] = (wval(num.find(w("abstractNumId"))), overrides)

    def style_attr(self, sid, key):
        seen = set()
        while sid in self.styles and sid not in seen:
            seen.add(sid)
            if self.styles[sid][key] is not None:
                return self.styles[sid][key]
            sid = self.styles[sid]["based"]
        return None

    def text(self, p):
        out = []

        def walk(el):
            for ch in el:
                tag = ch.tag
                if not isinstance(tag, str) or tag in DOCX_SKIP:
                    continue
                if tag == w("t"):
                    out.append(ch.text or "")
                elif tag == w("tab"):
                    out.append("\t")
                elif tag in (w("br"), w("cr")):
                    out.append("\n")
                elif tag == w("noBreakHyphen"):
                    out.append("-")
                elif tag == w("footnoteReference"):
                    out.append(f"[^{ch.get(w('id'))}]")
                elif tag == w("endnoteReference"):
                    out.append(f"[^e{ch.get(w('id'))}]")
                elif tag == w("fldChar"):
                    box = ch.find(f"{w('ffData')}/{w('checkBox')}")
                    if box is not None:
                        flag = box.find(w("checked"))
                        flag = box.find(w("default")) if flag is None else flag
                        on = flag is not None and wval(flag, "1") not in ("0", "false")
                        out.append("☒" if on else "☐")
                elif tag == w("hyperlink"):
                    start = len(out)
                    walk(ch)
                    url = self.links.get(ch.get(R_ID))
                    if url and collapse("".join(out[start:])) != url:
                        out.append(f" ({url})")
                elif tag == w("p"):  # a paragraph nested in a text box
                    if out and not out[-1].endswith("\n"):
                        out.append("\n")
                    walk(ch)
                else:
                    walk(ch)

        walk(p)
        return "".join(out)

    def heading(self, p):
        ppr = p.find(w("pPr"))
        sid = None
        if ppr is not None:
            level = to_int(wval(ppr.find(w("outlineLvl"))), -1)
            if 0 <= level < 9:
                return level + 1
            sid = wval(ppr.find(w("pStyle")))
        if not sid:
            return 0
        name = (self.styles.get(sid, {}).get("name") or sid).strip().lower()
        m = re.fullmatch(r"heading\s*([1-9])", name)
        if m:
            return int(m.group(1))
        if name == "title":
            return 1
        level = to_int(self.style_attr(sid, "outline"), -1)
        return level + 1 if 0 <= level < 9 else 0

    def label(self, p):
        """The list number Word would show (approximate), '-' for bullets, or None."""
        ppr = p.find(w("pPr"))
        num_id = ilvl = sid = None
        if ppr is not None:
            numpr = ppr.find(w("numPr"))
            if numpr is not None:
                num_id, ilvl = wval(numpr.find(w("numId"))), wval(numpr.find(w("ilvl")))
            sid = wval(ppr.find(w("pStyle")))
        if num_id is None and sid:
            num_id = self.style_attr(sid, "numId")
            ilvl = ilvl if ilvl is not None else self.style_attr(sid, "ilvl")
        if num_id in (None, "0") or num_id not in self.nums:
            return None
        aid, overrides = self.nums[num_id]
        levels = self.abstract.get(aid)
        if not levels:
            return None
        ilvl = to_int(ilvl)
        key = ("num", num_id) if overrides else ("abstract", aid)
        counters = self.counters.setdefault(key, {})
        if overrides and num_id not in self.started:
            self.started.add(num_id)
            counters.clear()
            counters.update({lvl: start - 1 for lvl, start in overrides.items()})
        fmt, pattern, start = levels.get(ilvl, ("decimal", f"%{ilvl + 1}.", 1))
        counters[ilvl] = counters.get(ilvl, start - 1) + 1
        for deeper in [k for k in counters if k > ilvl]:
            del counters[deeper]
        if fmt == "bullet":
            return "-"
        if fmt == "none":
            return None

        def sub(m):
            k = int(m.group(1)) - 1
            f, _, s = levels.get(k, ("decimal", "", 1))
            return list_number(counters.get(k, s), f)

        text = re.sub(r"%([1-9])", sub, pattern).strip()
        if text:
            self.numbered += 1
        return text or None

    def cell_text(self, tc):
        parts = []
        for kind, el in iter_blocks(tc):
            if kind == "p":
                label, text = self.label(el), self.text(el).strip()
                if text:
                    parts.append(f"{label} {text}" if label else text)
            else:
                for row in self.grid(el):
                    if any(row):
                        parts.append(" | ".join(row))
        return "\n".join(parts)

    def grid(self, tbl):
        """Rows of cell texts; merged cells repeat their value across the span."""
        rows, above = [], {}
        for tr in iter_tagged(tbl, w("tr")):
            values, col = {}, 0
            trpr = tr.find(w("trPr"))
            if trpr is not None:
                col += to_int(wval(trpr.find(w("gridBefore"))), 0)
            for tc in iter_tagged(tr, w("tc")):
                tcpr = tc.find(w("tcPr"))
                span, vmerge = 1, None
                if tcpr is not None:
                    span = max(1, to_int(wval(tcpr.find(w("gridSpan"))), 1))
                    vm = tcpr.find(w("vMerge"))
                    if vm is not None:
                        vmerge = wval(vm, "continue")
                text = above.get(col, "") if vmerge == "continue" else self.cell_text(tc)
                for k in range(span):
                    values[col + k] = text
                col += span
            rows.append([values.get(c, "") for c in range(max(values) + 1)] if values else [])
            above = values
        return rows


def list_number(n, fmt):
    if fmt in ("lowerLetter", "upperLetter"):
        s = chr(97 + (n - 1) % 26) * ((n - 1) // 26 + 1) if n > 0 else str(n)
        return s.upper() if fmt == "upperLetter" else s
    if fmt in ("lowerRoman", "upperRoman") and 0 < n < 4000:
        s = ""
        for value, numeral in ((1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"),
                               (90, "xc"), (50, "l"), (40, "xl"), (10, "x"), (9, "ix"),
                               (5, "v"), (4, "iv"), (1, "i")):
            while n >= value:
                s, n = s + numeral, n - value
        return s.upper() if fmt == "upperRoman" else s
    if fmt == "decimalZero":
        return f"{n:02d}"
    return str(n)


def docx_notes(doc, word):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from lxml import etree

    sections, counts = [], {"footnotes": 0, "endnotes": 0, "header_footer_lines": 0}
    rels = list(doc.part.rels.values())
    for title, reltype, item, mark, key in (
            ("Footnotes", RT.FOOTNOTES, "footnote", "", "footnotes"),
            ("Endnotes", RT.ENDNOTES, "endnote", "e", "endnotes")):
        lines = []
        for rel in rels:
            if rel.reltype != reltype or rel.is_external:
                continue
            root = etree.fromstring(rel.target_part.blob)
            for note in root.findall(w(item)):
                if note.get(w("type")) in ("separator", "continuationSeparator",
                                           "continuationNotice"):
                    continue
                text = " ".join(t for t in (word.text(p).strip() for kind, p in iter_blocks(note)
                                            if kind == "p") if t)
                if text:
                    lines.append(f"[^{mark}{note.get(w('id'))}] {one_line(text)}")
        if lines:
            sections.append(f"## {title}\n\n" + "\n".join(lines) + "\n")
            counts[key] = len(lines)
    seen = []
    for rel in rels:
        if rel.reltype in (RT.HEADER, RT.FOOTER) and not rel.is_external:
            root = etree.fromstring(rel.target_part.blob)
            for kind, el in iter_blocks(root):
                texts = ([word.text(el).strip()] if kind == "p"
                         else [" | ".join(r) for r in word.grid(el) if any(r)])
                for text in texts:
                    text = one_line(text)
                    if text and text not in seen:
                        seen.append(text)
    if seen:
        sections.append("## Headers and footers\n\n" + "\n".join(seen) + "\n")
        counts["header_footer_lines"] = len(seen)
    return "\n".join(sections), counts


def extract_docx(src, run, args):
    docx = need("docx", "docx")
    need("lxml.etree", "docx")
    try:
        doc = docx.Document(str(src))
    except ValueError as e:
        raise Fail(f"{src.name} is not a plain Word document ({e}).",
                   "Macro-enabled files and templates are not supported: save it as .docx "
                   "(Word: File > Save As > Word Document) and extract that copy.", 2)
    except Exception as e:
        raise Fail(f"Could not open {src.name} as a Word document ({type(e).__name__}: {e}).",
                   "Check that it opens in Word; re-save it as .docx (or PDF) and extract "
                   "that copy.", 1)
    word = Docx(doc)
    body = doc.element.body
    lines, n_par, shown, headings, labelled, n_tables = [], 0, 0, 0, 0, 0
    for kind, el in iter_blocks(body):
        if kind == "p":
            n_par += 1
            label = word.label(el)  # counted even for empty paragraphs, as Word does
            text = word.text(el).strip()
            if not text:
                continue
            shown += 1
            level = word.heading(el)
            headings += bool(level)
            labelled += bool(label)
            prefix = f"¶{n_par} " + ("#" * level + " " if level else "")
            lines.append(prefix + (f"{label} " if label else "") + one_line(text))
            continue
        n_tables += 1
        rel = f"table-{n_tables:02d}.csv"
        grid = word.grid(el)
        width = max((len(r) for r in grid), default=0)
        written = 0
        with run.csv(rel) as (out, entry):
            out.writerow(["_row"] + [col_letter(c) for c in range(1, width + 1)])
            for r, row in enumerate(grid, 1):
                if any(cell.strip() for cell in row):
                    out.writerow([r] + [one_line(c) for c in row] + [""] * (width - len(row)))
                    written += 1
            size = f"{plural(written, 'row')} x {plural(width, 'col')}"
            entry["what"] = f"table {n_tables} ({size}), after ¶{n_par}"
        lines.append(f"[{rel}: {size}]")
    run.text("text.md", "\n".join(lines) + "\n" if lines else "",
             f"{plural(n_par, 'body paragraph')} as ¶-numbered lines ({shown} with text) and "
             f"{plural(n_tables, 'table placeholder')}")
    numbered = word.numbered  # before the headers/footers advance the counters
    notes, counts = docx_notes(doc, word)
    if notes:
        run.text("notes.md", notes, "footnotes ([^N] in text.md), endnotes ([^eN]), and header "
                                    "and footer text")
    if not lines:
        run.warn("the document body contains no text.")
    if numbered:
        run.warn(f"{plural(numbered, 'paragraph')} {'uses' if numbered == 1 else 'use'} Word's "
                 f"automatic list numbering; "
                 f"those numbers (e.g. '3.') were reconstructed and can differ from what Word "
                 f"shows, so cite the ¶ number.")
    if body.find(f".//{w('ins')}") is not None or body.find(f".//{w('del')}") is not None:
        run.warn("the document has tracked changes; text.md shows them accepted (insertions "
                 "kept, deletions dropped).")
    run.stats.update({"engine": f"python-docx {version_of(docx)}", "paragraphs": n_par,
                      "paragraphs_with_text": shown, "headings": headings,
                      "list_items": labelled, "tables": n_tables, **counts})
    return LOCATORS["docx"]


# --------------------------------------------------------------------------- HTML

HTML_SKIP = {"script", "style", "noscript", "template", "svg"}
HTML_HEADINGS = {"h1", "h2", "h3", "h4", "h5", "h6"}
HTML_BLOCKS = {"address", "article", "aside", "blockquote", "body", "caption", "center", "dd",
               "details", "dialog", "div", "dt", "fieldset", "figcaption", "figure", "footer",
               "form", "header", "hgroup", "hr", "html", "legend", "main", "nav", "option", "p",
               "section", "summary"}
CELL_BREAKS = HTML_BLOCKS | HTML_HEADINGS | {"br", "li", "ul", "ol", "dl", "pre"}
BREAK = "\x00"  # line-break marker inside a table cell


class HtmlText(HTMLParser):
    """HTML -> markdown-ish lines: # headings, - items, ' | ' table rows, text (url)."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.lines, self.buf, self.prefix, self.skip = [], [], "", []
        self.title, self.in_title, self.pre, self.lists = [], False, 0, 0
        self.depth, self.row, self.cell, self.anchors = 0, None, None, []
        self.links = self.headings = self.tables = 0

    def sink(self):
        return self.cell if self.cell is not None else self.buf

    def gap(self):
        if self.lines and self.lines[-1] != "":
            self.lines.append("")

    def flush(self):
        if self.cell is not None:
            return
        text = collapse("".join(self.buf))
        self.buf = []
        if text:
            self.lines.append(self.prefix + text)
            self.prefix = " " * len(self.prefix)  # later lines of the same item or heading

    def flush_pre(self):
        text = "".join(self.buf).rstrip()
        self.buf = []
        if text.strip():
            self.lines.append(self.prefix + text)
        else:
            self.gap()

    def end_cell(self):
        if self.cell is not None:
            parts = [collapse(p) for p in "".join(self.cell).split(BREAK)]
            self.row.append("\\n".join(p for p in parts if p))
            self.cell = None

    def end_row(self):
        self.end_cell()
        if self.row is not None and any(self.row):
            self.lines.append(" | ".join(self.row))
        self.row = None

    def handle_starttag(self, tag, attrs):
        if self.in_title and tag != "title":  # title holds text only; a tag means it ended
            self.in_title = False
        if self.skip or tag in HTML_SKIP:
            if tag in HTML_SKIP:
                self.skip.append(tag)
            return
        if tag == "title":
            self.in_title = True
        elif tag == "a":
            sink = self.sink()
            self.anchors.append(((dict(attrs).get("href") or "").strip(), sink, len(sink)))
        elif tag == "table":
            if self.depth == 0:
                self.flush()
                self.gap()
                self.tables += 1
            else:
                self.sink().append(" ")
            self.depth += 1
        elif self.depth > 1 and tag in ("tr", "td", "th"):  # nested table: flatten into cell
            self.sink().append(BREAK if tag == "tr" else " ")
        elif self.depth == 1 and tag == "tr":
            self.end_row()
            self.flush()
            self.row = []
        elif self.depth == 1 and tag in ("td", "th"):
            if self.row is None:
                self.row = []
            self.end_cell()
            self.cell = []
        elif self.cell is not None:
            if tag in CELL_BREAKS:
                self.cell.append(BREAK)
        elif tag in HTML_HEADINGS:
            self.flush()
            self.gap()
            self.prefix = "#" * int(tag[1]) + " "
            self.headings += 1
        elif tag == "li":
            self.flush()
            self.prefix = "  " * max(0, self.lists - 1) + "- "
        elif tag in ("ul", "ol", "dl"):
            self.flush()
            self.lists += 1
        elif tag == "br":
            self.flush_pre() if self.pre else self.flush()
        elif tag == "pre":
            self.flush()
            self.pre += 1
        elif tag in HTML_BLOCKS:
            self.flush()

    def handle_endtag(self, tag):
        if self.skip:
            if tag in self.skip:
                while self.skip and self.skip.pop() != tag:
                    pass
            return
        if tag == "title":
            self.in_title = False
        elif tag == "a":
            self.end_link()
        elif tag == "table":
            if self.depth == 1:
                self.end_row()
                self.flush()
                self.gap()
            elif self.depth > 1:
                self.sink().append(BREAK)
            self.depth = max(0, self.depth - 1)
        elif self.depth > 1 and tag in ("tr", "td", "th"):
            pass
        elif self.depth == 1 and tag in ("td", "th"):
            self.end_cell()
        elif self.depth == 1 and tag == "tr":
            self.end_row()
        elif self.cell is not None:
            if tag in CELL_BREAKS:
                self.cell.append(BREAK)
        elif tag in HTML_HEADINGS or tag == "li":
            self.flush()
            self.prefix = ""
        elif tag in ("ul", "ol", "dl"):
            self.flush()
            self.lists = max(0, self.lists - 1)
            self.prefix = ""
        elif tag == "pre":
            self.flush_pre()
            self.pre = max(0, self.pre - 1)
        elif tag in HTML_BLOCKS:
            self.flush()

    def end_link(self):
        if not self.anchors:
            return
        href, sink, start = self.anchors.pop()
        if not href or href.startswith("#") or href.lower().startswith("javascript:"):
            return
        self.links += 1
        text = collapse("".join(sink[start:])) if self.sink() is sink else ""
        if text != href:
            self.sink().append(f" ({href})")

    def handle_data(self, data):
        if self.skip:
            return
        if self.in_title:
            self.title.append(data)
            return
        data = data.replace(BREAK, "")
        if self.pre and self.cell is None:
            parts = data.split("\n")
            self.buf.append(parts[0])
            for part in parts[1:]:
                self.flush_pre()
                self.buf.append(part)
            return
        self.sink().append(data)

    def result(self):
        self.close()
        self.end_row()
        self.flush_pre() if self.pre else self.flush()
        while self.lines and self.lines[-1] == "":
            self.lines.pop()
        return "\n".join(self.lines) + "\n" if self.lines else ""


def decode_html(data):
    for bom, enc in ((codecs.BOM_UTF8, "utf-8-sig"), (codecs.BOM_UTF16_LE, "utf-16"),
                     (codecs.BOM_UTF16_BE, "utf-16")):
        if data.startswith(bom):
            return data.decode(enc, "replace"), enc
    m = re.search(rb"""<meta[^>]+charset\s*=\s*["']?\s*([A-Za-z0-9_.:-]+)""", data[:16384], re.I)
    if m:
        declared = m.group(1).decode("ascii", "ignore").lower()
        try:
            codecs.lookup(declared)
            return data.decode(declared, "replace"), declared
        except LookupError:
            pass
    return decode_text(data)


def parse_html(text):
    parser = HtmlText()
    parser.feed(text)
    return parser, parser.result()


def extract_html(src, run, args):
    text, encoding = decode_html(src.read_bytes())
    text = norm_newlines(text)
    parser, out = parse_html(text)
    if parser.in_title:  # HTML5 parsers read an unclosed <title> to the end of the file
        run.warn("the page's <title> is never closed; it was dropped and the page re-read.")
        parser, out = parse_html(re.sub(r"(?i)<title\b[^>]*>", "", text, count=1))
    lines = out.count("\n")
    run.text("text.md", out, f"the page text in reading order, {plural(lines, 'line')}")
    if not out.strip():
        run.warn("the page contains no text (it may be built by JavaScript); open it in a "
                 "browser and save the rendered page, or print it to PDF.")
    run.stats.update({"encoding": encoding, "title": collapse("".join(parser.title))[:200] or None,
                      "chars": len(out), "lines": lines, "links": parser.links,
                      "headings": parser.headings, "tables": parser.tables})
    return LOCATORS["html"]


# ----------------------------------------------------------------------- RTF, JSON, XML

def extract_rtf(src, run, args):
    text = converter = None
    tool = shutil.which("textutil")
    if tool:
        try:
            res = subprocess.run([tool, "-convert", "txt", "-encoding", "UTF-8", "-stdout",
                                  str(src.resolve())], capture_output=True, timeout=300)
            if res.returncode == 0:
                text, converter = res.stdout.decode("utf-8", "replace"), "textutil"
            else:
                reason = one_line(res.stderr.decode("utf-8", "replace").strip())[:200]
                run.warn(f"textutil could not convert the file ({reason}); used striprtf.")
        except (OSError, subprocess.TimeoutExpired) as e:
            run.warn(f"textutil failed ({type(e).__name__}); used striprtf.")
    if text is None:
        module = need("striprtf.striprtf", "rtf")
        raw = src.read_bytes().decode("latin-1")  # RTF is 7-bit; latin-1 keeps every byte
        m = re.search(r"\\ansicpg(\d+)", raw[:4096])
        encoding = f"cp{m.group(1)}" if m else "cp1252"
        try:
            codecs.lookup(encoding)
        except LookupError:
            encoding = "cp1252"
        try:
            text = module.rtf_to_text(raw, encoding=encoding, errors="replace")
        except TypeError:  # older striprtf without these keywords
            text = module.rtf_to_text(raw)
        converter = "striprtf"
    text = norm_newlines(text).replace("\u2028", "\n").replace("\u2029", "\n")
    text = text.replace("\x00", "")
    if text and not text.endswith("\n"):
        text += "\n"
    lines = text.count("\n")
    run.text("text.md", text, f"the document as plain text ({converter}), "
                              f"{plural(lines, 'line')}")
    if not text.strip():
        run.warn("the document contains no text.")
    run.stats.update({"converter": converter, "lines": lines, "chars": len(text)})
    return LOCATORS["rtf"]


def json_type(value):
    return {dict: "object", list: "array", str: "string", bool: "boolean",
            type(None): "null"}.get(type(value), "number")


def extract_json(src, run, args):
    data = src.read_bytes()
    try:
        value = json.loads(data.decode("utf-8-sig"))
    except UnicodeDecodeError:
        try:
            value = json.loads(data)  # detects UTF-16/32
        except (ValueError, RecursionError) as e:
            run.warn(f"not valid JSON ({e}); extracted as plain text instead.")
            return extract_plain(src, run, args)
    except (ValueError, RecursionError) as e:
        run.warn(f"not valid JSON ({e}); extracted as plain text instead.")
        return extract_plain(src, run, args)
    pretty = json.dumps(value, indent=2, ensure_ascii=False) + "\n"
    lines = pretty.count("\n")
    run.text("pretty.json", pretty, f"the JSON pretty-printed, {plural(lines, 'line')}")
    stats = {"type": json_type(value), "lines": lines}
    if isinstance(value, dict):
        stats.update({"keys": [k[:100] for k in list(value)[:100]], "key_count": len(value)})
    elif isinstance(value, list):
        stats["length"] = len(value)
        keys = {}
        for item in value[:100]:
            if isinstance(item, dict):
                keys.update(dict.fromkeys(item))
        if keys:
            stats["item_keys"] = [k[:100] for k in list(keys)[:100]]
    run.stats.update(stats)
    return LOCATORS["json"]


def extract_xml(src, run, args):
    import xml.etree.ElementTree as ET
    from xml.dom import minidom

    count, root = 0, None
    try:
        for event, el in ET.iterparse(str(src), events=("start", "end")):
            if event == "start":
                count += 1
                root = root or el.tag
            else:
                el.clear()
    except ET.ParseError as e:
        run.warn(f"not well-formed XML ({e}); extracted as plain text instead.")
        return extract_plain(src, run, args)
    namespace, _, local = root[1:].rpartition("}") if root.startswith("{") else ("", "", root)
    run.stats.update({"root": local, "namespace": namespace or None, "elements": count})
    if run.size > XML_PRETTY_MAX:
        text, _ = decode_text(src.read_bytes())
        name, how = "source.xml", "the XML as-is (too large to pretty-print)"
        run.warn(f"the XML is over {XML_PRETTY_MAX >> 20} MB, so it was copied as-is to "
                 f"source.xml instead of being pretty-printed.")
    else:
        dom = minidom.parse(str(src))
        stack = [dom.documentElement]
        while stack:  # drop whitespace-only text so the indentation is minidom's alone
            node = stack.pop()
            for child in list(node.childNodes):
                if child.nodeType == child.TEXT_NODE and not child.data.strip():
                    node.removeChild(child)
                elif child.nodeType == child.ELEMENT_NODE:
                    stack.append(child)
        text = dom.toprettyxml(indent="  ", encoding="utf-8").decode("utf-8")
        name, how = "pretty.xml", "the XML pretty-printed"
        dom.unlink()
    text = norm_newlines(text)
    if text and not text.endswith("\n"):
        text += "\n"
    lines = text.count("\n")
    run.text(name, text, f"{how}, {plural(lines, 'line')}")
    run.stats.update({"pretty": name == "pretty.xml", "lines": lines})
    return LOCATORS["xml"].format(file=name)


EXTRACTORS = {
    "xlsx": extract_xlsx, "xls": extract_xls, "pdf": extract_pdf, "docx": extract_docx,
    "html": extract_html, "rtf": extract_rtf, "json": extract_json, "xml": extract_xml,
    "csv": lambda s, r, a: extract_delimited(s, r, a, "csv"),
    "tsv": lambda s, r, a: extract_delimited(s, r, a, "tsv"),
    "txt": lambda s, r, a: extract_plain(s, r, a, table_ok=True),
    "md": lambda s, r, a: extract_plain(s, r, a),
    "jsonl": lambda s, r, a: extract_plain(s, r, a),      # one JSON value per line: cited by line, as text
    "code": lambda s, r, a: extract_code(s, r, a),
}
DEPS = {"xlsx": ["openpyxl"], "xls": ["xlrd"], "pdf": ["pypdf"], "docx": ["docx", "lxml.etree"]}


# --------------------------------------------------------------------------- main

def write_manifest(out, manifest):
    tmp = out / "manifest.json.tmp"
    tmp.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    os.replace(tmp, out / "manifest.json")


def capped(value):
    """Cut every list longer than STDOUT_CAP, adding a sibling <key>_total."""
    if isinstance(value, list):
        return [capped(v) for v in value]
    if not isinstance(value, dict):
        return value
    out = {}
    for key, v in value.items():
        if isinstance(v, list) and len(v) > STDOUT_CAP:
            out[key] = [capped(x) for x in v[:STDOUT_CAP]]
            out[f"{key}_total"] = value.get(f"{key}_total", len(v))
        elif key not in out:
            out[key] = capped(v)
    return out


def execute(argv):
    args = build_parser().parse_args(argv)
    if args.max_pages is not None and args.max_pages < 1:
        raise Fail("--max-pages must be a positive whole number.",
                   "Omit --max-pages to extract every page, or pass e.g. --max-pages 50.", 1)
    src = Path(args.source)
    if not src.exists():
        raise Fail(f"Source not found: {args.source}",
                   "Check the path; a relative path resolves from the current directory.", 1)
    if not src.is_file():
        raise Fail(f"Source is not a file: {args.source}",
                   "Pass one document file per run, not a directory.", 1)
    ext = src.suffix.lower()
    if ext in DATA_EXTS:
        raise Fail(DATA_ERROR, DATA_HINT, 2)
    size = src.stat().st_size
    if size == 0:
        raise Fail(f"Source is empty (0 bytes): {args.source}",
                   "Re-download or re-export the document, then extract it again.", 1)
    if size > MAX_BYTES:
        raise Fail(f"Source is {size / 2**20:.0f} MB; the limit is {MAX_BYTES >> 20} MB.",
                   "Metadata documents are rarely this large: make sure it is not a data file, "
                   "or extract a smaller export (e.g. just the codebook section).", 2)
    try:
        with open(src, "rb") as f:
            head = f.read(8192)
    except OSError as e:
        raise Fail(f"Cannot read {args.source}: {e.strerror or e}",
                   "Check the file's permissions and that it is fully downloaded or synced.", 1)
    signature = data_signature(head)
    if signature:
        raise Fail(f"{DATA_ERROR} (Its content is a {signature} file.)", DATA_HINT, 2)
    kind, note = detect_kind(src, ext, head)
    if not args.documentation and (kind in ("csv", "tsv") or (
            kind == "txt" and consistent_delimiter(decode_text(head)[0]))):
        raise Fail("A delimited file can be a dataset as easily as a data dictionary, and the "
                   "agent must not open data files.",
                   "Ask the user whether this file describes variables (one row per variable or "
                   "code) or holds records (one row per person, visit, or sample). Only if it "
                   "describes variables, re-run with --documentation.", 2)
    if not args.documentation and kind in ("json", "jsonl"):
        raise Fail("A JSON file can be a dataset (an array of records, or one record per line) or "
                   "documentation (a JSON Schema, a REDCap metadata export); the script cannot "
                   "tell them apart, and the agent must not open data files.",
                   "Ask the user whether this file describes variables (a schema, a codebook or "
                   "metadata export: one entry per variable or code) or holds records (one entry "
                   "per person, visit, or sample). Only if it describes variables, re-run with "
                   "--documentation.", 2)

    out = Path(args.out)
    skill = schemify_skill_above(out)
    if skill:
        raise Fail(f"Refusing to write inside the schemify skill directory ({skill}).",
                   "Pass an --out directory in the steward's workspace (e.g. beside the "
                   "package's sources), never inside the installed skill.", 1)
    if out.exists() and not out.is_dir():
        raise Fail(f"--out exists and is not a directory: {args.out}",
                   "Pass a directory path; it is created if missing.", 1)
    prev = previous_manifest(out)
    with contextlib.suppress(ValueError):
        inside = src.resolve().relative_to(out.resolve()).as_posix()
        if inside == "manifest.json" or GENERATED.match(inside):
            raise Fail(f"The source {args.source} is one of extract.py's own output files.",
                       "Extract the original document, into a separate --out directory.", 1)
    for module in DEPS.get(kind, []):
        need(module, kind)
    if kind == "rtf" and not shutil.which("textutil"):
        need("striprtf.striprtf", kind)
    digest = sha256_file(src)
    created = not out.exists()
    try:
        out.mkdir(parents=True, exist_ok=True)
        clear_previous(out, prev)
    except OSError as e:
        raise Fail(f"Cannot prepare the output directory {args.out}: {e.strerror or e}",
                   "Check that the location is writable, or pass another --out.", 1)

    job = Run(out, size)
    if note:
        job.warn(note)
    if prev and prev.get("sha256") != digest:
        job.warn(f"{args.out} held an extraction of a different source "
                 f"({prev.get('source')}); its files were replaced.")
    if args.max_pages and kind != "pdf":
        job.warn("--max-pages applies to PDFs only; ignored.")
    base = {"ok": True, "source": args.source, "id": args.id, "kind": kind, "sha256": digest,
            "bytes": size, "extracted_at": now_iso()}
    try:
        with contextlib.redirect_stdout(sys.stderr), quiet_libraries(job):
            locators = EXTRACTORS[kind](src, job, args)
    except (Exception, KeyboardInterrupt) as e:
        failure = e if isinstance(e, Fail) else Fail(
            "Interrupted." if isinstance(e, KeyboardInterrupt) else
            f"Extraction failed ({kind}): {type(e).__name__}: {e}",
            "The file may be damaged or unusual: open it in its own application, re-save or "
            "export it (e.g. as .xlsx, .docx or PDF), then extract that copy.", 1)
        with contextlib.suppress(OSError):
            if job.files:  # a manifest of the partial output lets the next run clean it up
                write_manifest(out, {**base, "ok": False, "error": failure.error,
                                     "hint": failure.hint, "locators": "", "files": job.files,
                                     "warnings": job.warnings, "stats": job.stats})
            elif created:
                out.rmdir()
        raise failure from None
    manifest = {**base, "locators": locators, "files": job.files, "files_total": len(job.files), "warnings": job.warnings,
                "stats": job.stats}
    write_manifest(out, manifest)
    return manifest


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):
        with contextlib.suppress(AttributeError, ValueError, io.UnsupportedOperation):
            stream.reconfigure(encoding="utf-8")
    try:
        manifest = execute(argv)
    except Fail as f:
        print(json.dumps({"ok": False, "error": f.error, "hint": f.hint}, indent=2,
                         ensure_ascii=False))
        return f.code
    except Exception as e:  # never a traceback
        print(json.dumps({"ok": False, "error": f"Unexpected error: {type(e).__name__}: {e}",
                          "hint": "Re-run once; if it fails again, convert the document (e.g. "
                                  "to PDF or .xlsx) and extract that copy."}, indent=2))
        return 1
    print(json.dumps(capped(manifest), indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
